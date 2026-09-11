# -*- coding: utf-8 -*-
"""Modul Bonusy AO — 14 pobočkových dlaždic, v každé sekce „Podkladová tabulka“.

Struktura podkladové tabulky je 1:1 podle vzoru
„vzorová podkladová tabulka.xlsx“ (list Podklad_*): 47 sloupců od
„IČO - zadané“ po „nastavení Sklad6“.
"""

import asyncio
import datetime
import inspect
import io
import os
import re
import unicodedata

from nicegui import app, ui

import intranet_data
import intranet_logger
from intranet_ui_utils import refreshable_na_klienta

POBOCKY = [
    'Praha', 'Pardubice', 'Most', 'Jilemnice', 'Liberec', 'Nová role', 'Brno',
    'Ostrava', 'Plzeň', 'Olomouc', 'České Budějovice', 'Hodonín',
    'Horšovský Týn', 'Zlín',
]

INTERVALY = ['Měsíční', 'Kvartální', 'Pololetní', 'Roční']
SKLAD6 = ['Ano', 'Ano – Centrála', 'Ano – Skupinový']
AN = ['A', 'N']

# (db pole, hlavička ve vzoru, typ, šířka v gridu)
# typ: text | an | datum | cislo | procento | interval | sklad6
SLOUPCE = [
    ('ico_zadane',             'IČO - zadané',                            'text',     110),
    ('pobocka',                'Pobočka',                                 'text',     110),
    ('databaze',               'Databáze',                                'text',      90),
    ('ico_kontrola',           'IČO - kontrola',                          'text',     110),
    ('nazev_jmeno',            'Název jméno',                             'text',     200),
    ('nazev_provozovny',       'Název provozovny',                        'text',     200),
    ('dalsi_odbery',           'Další odběry (pobočky)',                  'text',     160),
    ('typ_bonusu',             'Typ bonusu',                              'text',     120),
    ('interval',               'Interval',                                'interval', 120),
    ('datum_od',               'Od',                                      'datum',    105),
    ('datum_do',               'Do',                                      'datum',    105),
    ('majetek',                'Majetek',                                 'text',     130),
    ('investice',              'Investice',                               'text',     130),
    ('cislo_dokladu',          'Číslo dokladu s prodlouženou splatností', 'text',     170),
    ('obrat_dp',               'Obrat - DP',                              'an',        95),
    ('bonus_dp',               'Bonus - DP',                              'an',        95),
    ('obrat_akce',             'Obrat - Akce',                            'an',       100),
    ('bonus_akce',             'Bonus - Akce',                            'an',       100),
    ('tabak',                  'Tabák',                                   'an',        80),
    ('ceniny',                 'Ceniny',                                  'an',        80),
    ('nektar_natura',          'Nektar Natura',                           'an',       110),
    ('premier_wines',          'Premier Wines',                           'an',       110),
    ('eso',                    'ESO',                                     'an',        70),
    ('pivo',                   'Pivo',                                    'an',        70),
    ('vyjmuto_z_obratu',       'Vyjmuto z obratu',                        'text',     160),
    ('bonus_pouze_nadrazena',  'Bonus pouze z nadřazené skupiny',         'text',     200),
    ('minimalni_odber',        'Minimální odběr',                         'cislo',    120),
    ('skupina',                'Skupina',                                 'text',     140),
    ('poznamka',               'Poznámka',                                'text',     200),
    ('poznamka_investice',     'Poznámka k investici',                    'text',     200),
    ('procento_bonusu',        '% bonusu',                                'procento',  95),
    ('procento_bonusu_tabak',  '% bonusu tabák',                          'procento', 110),
] + [
    sl
    for i in range(1, 8)
    for sl in ((f'interval_bonusu_{i}', 'interval bonusu', 'cislo', 120),
               (f'procento_{i}',        '%',               'procento', 80))
] + [
    ('sklad6',                 'nastavení Sklad6',                        'sklad6',   150),
]

_SQL_TYP = {
    'text':     'TEXT',
    'an':       'VARCHAR(3)',
    'datum':    'DATE',
    'cislo':    'DECIMAL(18,2)',
    'procento': 'DECIMAL(9,6)',
    'interval': 'VARCHAR(30)',
    'sklad6':   'VARCHAR(30)',
}
_TYP = {f: t for f, _h, t, _w in SLOUPCE}
POLE = [f for f, _h, _t, _w in SLOUPCE]
# `interval` je v MySQL rezervované slovo → všechna pole v SQL v backticích.
POLE_SQL = ', '.join(f'`{f}`' for f in POLE)

TABULKA = 'bonusy_ao_podklad'


# ─── normalizace / převody ────────────────────────────────────────────────────

def _norm(s) -> str:
    """Malá písmena bez diakritiky, sloučené mezery — pro porovnávání hlaviček."""
    s = ''.join(c for c in unicodedata.normalize('NFKD', str(s or ''))
                if not unicodedata.combining(c))
    return re.sub(r'\s+', ' ', s).strip().lower()


_INTERVAL_MAPA = {
    'mesicne': 'Měsíční', 'mesicni': 'Měsíční',
    'ctvrtletne': 'Kvartální', 'kvartalni': 'Kvartální', 'ctvrtletni': 'Kvartální',
    'pololetne': 'Pololetní', 'pololetni': 'Pololetní',
    'rocne': 'Roční', 'rocni': 'Roční',
}
_SKLAD6_MAPA = {
    'ano': 'Ano',
    'ano - centrala': 'Ano – Centrála', 'ano - na centrale': 'Ano – Centrála',
    'ano centrala': 'Ano – Centrála',
    'ano - skupinovy': 'Ano – Skupinový', 'ano skupinovy': 'Ano – Skupinový',
}


def preved(typ: str, hodnota):
    """Hodnota z Excelu / z editace gridu → hodnota do DB (nebo None)."""
    if hodnota is None:
        return None
    if isinstance(hodnota, str):
        hodnota = hodnota.strip()
        if not hodnota:
            return None

    if typ == 'datum':
        if isinstance(hodnota, datetime.datetime):
            return hodnota.date()
        if isinstance(hodnota, datetime.date):
            return hodnota
        for fmt in ('%d.%m.%Y', '%Y-%m-%d', '%d.%m.%y', '%d/%m/%Y'):
            try:
                return datetime.datetime.strptime(str(hodnota).strip(), fmt).date()
            except ValueError:
                continue
        return None

    if typ in ('cislo', 'procento'):
        if isinstance(hodnota, bool):
            return None
        if isinstance(hodnota, (int, float)):
            cislo = float(hodnota)
        else:
            txt = re.sub(r'[^\d,.\-]', '', str(hodnota)).replace(',', '.')
            if txt.count('.') > 1:            # 1.234.567 → oddělovač tisíců
                txt = txt.replace('.', '')
            try:
                cislo = float(txt)
            except ValueError:
                return None
            # ponytail: procento zadané jako „3" bereme jako 3 %, „0,03" jako zlomek.
            # Hranice 1 = 100 %; vyšší bonus než 100 % neexistuje.
            if typ == 'procento' and cislo > 1:
                cislo /= 100.0
        return round(cislo, 6 if typ == 'procento' else 2)

    if typ == 'an':
        z = _norm(hodnota)[:1].upper()
        return z if z in ('A', 'N') else None

    if typ == 'interval':
        return _INTERVAL_MAPA.get(_norm(hodnota)) or (
            hodnota if hodnota in INTERVALY else None)

    if typ == 'sklad6':
        return _SKLAD6_MAPA.get(_norm(hodnota).replace('–', '-')) or (
            hodnota if hodnota in SKLAD6 else None)

    return str(hodnota)[:2000]


def mapuj_hlavicky(hlavicka: list) -> dict:
    """Hlavičkový řádek z Excelu → {index sloupce: db pole}.

    Opakující se hlavičky („interval bonusu", „%") se párují zleva doprava,
    takže sedmice dvojic skončí ve správných polích.
    """
    volne = list(enumerate(_norm(h) for h in hlavicka))
    mapa = {}
    for pole, hlav, _t, _w in SLOUPCE:
        cil = _norm(hlav)
        for poz, (idx, h) in enumerate(volne):
            if h == cil:
                mapa[idx] = pole
                del volne[poz]
                break
    return mapa


# ─── DB ───────────────────────────────────────────────────────────────────────

def inicializace_bonusy_ao_db():
    conn = intranet_data.get_db_connection()
    if not conn:
        return
    try:
        sloupce_ddl = ',\n                '.join(
            f'`{f}` {_SQL_TYP[t]}' for f, _h, t, _w in SLOUPCE)
        cur = conn.cursor()
        cur.execute(f"""
            CREATE TABLE IF NOT EXISTS {TABULKA} (
                id      INT AUTO_INCREMENT PRIMARY KEY,
                pobocka_klic VARCHAR(50) NOT NULL,
                poradi  INT NOT NULL DEFAULT 0,
                {sloupce_ddl},
                INDEX idx_pobocka (pobocka_klic, poradi)
            ) CHARACTER SET utf8mb4 COLLATE utf8mb4_unicode_ci
        """)
        conn.commit()
        cur.close()
    except Exception as e:
        print(f'[bonusy_ao] Chyba inicializace DB: {e}')
    finally:
        conn.close()


def nacti_radky(pobocka: str) -> list[dict]:
    conn = intranet_data.get_db_connection()
    if not conn:
        return []
    try:
        cur = conn.cursor(dictionary=True)
        cur.execute(
            f'SELECT id, {POLE_SQL} FROM {TABULKA} '
            f'WHERE pobocka_klic=%s ORDER BY poradi, id', (pobocka,))
        radky = cur.fetchall()
        cur.close()
    finally:
        conn.close()
    for r in radky:
        for f in POLE:
            v = r[f]
            if _TYP[f] == 'datum':
                r[f] = v.strftime('%d.%m.%Y') if v else ''
            elif _TYP[f] in ('cislo', 'procento'):
                r[f] = float(v) if v is not None else None
            else:
                r[f] = v or ''
    return radky


def pridej_radek(pobocka: str) -> int:
    conn = intranet_data.get_db_connection()
    if not conn:
        return 0
    try:
        cur = conn.cursor()
        cur.execute(f'SELECT COALESCE(MAX(poradi), 0) + 1 FROM {TABULKA} '
                    f'WHERE pobocka_klic=%s', (pobocka,))
        poradi = cur.fetchone()[0]
        cur.execute(f'INSERT INTO {TABULKA} (pobocka_klic, poradi, `pobocka`) '
                    f'VALUES (%s, %s, %s)', (pobocka, poradi, pobocka))
        conn.commit()
        rid = cur.lastrowid
        cur.close()
        return rid
    finally:
        conn.close()


def uloz_bunku(rid: int, pole: str, hodnota) -> None:
    if pole not in POLE:                       # název sloupce jde do SQL přímo
        return
    conn = intranet_data.get_db_connection()
    if not conn:
        return
    try:
        cur = conn.cursor()
        cur.execute(f'UPDATE {TABULKA} SET `{pole}`=%s WHERE id=%s',
                    (preved(_TYP[pole], hodnota), rid))
        conn.commit()
        cur.close()
    finally:
        conn.close()


def smaz_radky(ids: list[int]) -> int:
    if not ids:
        return 0
    conn = intranet_data.get_db_connection()
    if not conn:
        return 0
    try:
        cur = conn.cursor()
        cur.execute(f'DELETE FROM {TABULKA} WHERE id IN '
                    f'({",".join(["%s"] * len(ids))})', tuple(ids))
        conn.commit()
        pocet = cur.rowcount
        cur.close()
        return pocet
    finally:
        conn.close()


def uloz_import(pobocka: str, radky: list[dict], nahradit: bool) -> int:
    conn = intranet_data.get_db_connection()
    if not conn:
        return 0
    try:
        cur = conn.cursor()
        if nahradit:
            cur.execute(f'DELETE FROM {TABULKA} WHERE pobocka_klic=%s', (pobocka,))
            poradi = 0
        else:
            cur.execute(f'SELECT COALESCE(MAX(poradi), 0) FROM {TABULKA} '
                        f'WHERE pobocka_klic=%s', (pobocka,))
            poradi = cur.fetchone()[0]
        sql = (f'INSERT INTO {TABULKA} (pobocka_klic, poradi, {POLE_SQL}) '
               f'VALUES ({", ".join(["%s"] * (len(POLE) + 2))})')
        data = []
        for r in radky:
            poradi += 1
            data.append((pobocka, poradi, *[r.get(f) for f in POLE]))
        cur.executemany(sql, data)
        conn.commit()
        cur.close()
        return len(data)
    finally:
        conn.close()


# ─── import z Excelu ──────────────────────────────────────────────────────────

def nacti_listy(raw: bytes) -> list[str]:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def parsuj_excel(raw: bytes, list_nazev: str, pobocka: str) -> tuple[list[dict], str]:
    """Vrátí (řádky připravené k uložení, chybová hláška)."""
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    try:
        ws = wb[list_nazev]
        radky_it = ws.iter_rows(values_only=True)
        try:
            hlavicka = next(radky_it)
        except StopIteration:
            return [], 'List je prázdný.'
        mapa = mapuj_hlavicky(list(hlavicka))
        if len(mapa) < 10:
            return [], ('V prvním řádku listu nejsou hlavičky podkladové tabulky '
                        f'(rozpoznáno {len(mapa)} sloupců).')
        vysledek = []
        for radek in radky_it:
            if not any(v is not None and str(v).strip() != '' for v in radek):
                continue
            zaznam = {f: None for f in POLE}
            for idx, pole in mapa.items():
                if idx < len(radek):
                    zaznam[pole] = preved(_TYP[pole], radek[idx])
            zaznam['pobocka'] = zaznam.get('pobocka') or pobocka
            vysledek.append(zaznam)
        return vysledek, ''
    except KeyError:
        return [], f'List „{list_nazev}" v souboru není.'
    finally:
        wb.close()


# ─── UI ───────────────────────────────────────────────────────────────────────

# Chování gridu 1:1 jako tabulky v modulu Sankce: výška do patičky okna a
# sloupce nejdřív roztažené na obsah, pak dorovnané na plnou šířku tabulky.
_GRID_STYLE = 'height: calc(100vh - 360px); min-height: 420px'
_BEZ_TYPU = '(bez typu)'


def _typ_bonusu(radek: dict) -> str:
    return str(radek.get('typ_bonusu') or '').strip() or _BEZ_TYPU


_AUTOSIZE_FIT = (
    "function(p){var api=p.api;if(!api||!api.autoSizeAllColumns||!api.sizeColumnsToFit)return;"
    "requestAnimationFrame(function(){api.autoSizeAllColumns();"
    "var cols=api.getColumns?api.getColumns():null;if(!cols)return;"
    "var lim=cols.filter(function(c){return c.isVisible();})"
    ".map(function(c){return{key:c.getColId(),minWidth:Math.ceil(c.getActualWidth())};});"
    "api.sizeColumnsToFit({columnLimits:lim});});}"
)

_FMT_PROCENTO = ("function(p){return (p.value===null||p.value===undefined||p.value==='')"
                 " ? '' : (p.value*100).toFixed(2).replace('.',',')+' %';}")
_FMT_CISLO = ("function(p){return (p.value===null||p.value===undefined||p.value==='')"
              " ? '' : Number(p.value).toLocaleString('cs-CZ');}")


def _col_defs() -> list[dict]:
    cols = [{
        'headerName': '', 'field': '_sel', 'width': 44, 'minWidth': 44, 'maxWidth': 44,
        'pinned': 'left', 'checkboxSelection': True, 'headerCheckboxSelection': True,
        'sortable': False, 'editable': False, 'resizable': False, 'filter': False,
        'suppressMovable': True, 'headerTooltip': 'Označení řádků ke smazání',
    }]
    for pole, hlav, typ, sirka in SLOUPCE:
        col = {'headerName': hlav, 'field': pole, 'width': sirka,
               'editable': pole != 'pobocka', 'sortable': True, 'filter': True,
               'headerTooltip': hlav}
        if typ in ('interval', 'sklad6', 'an'):
            col['cellEditor'] = 'agSelectCellEditor'
            col['cellEditorParams'] = {
                'values': {'interval': INTERVALY, 'sklad6': SKLAD6, 'an': AN}[typ]}
        elif typ == 'procento':
            col[':valueFormatter'] = _FMT_PROCENTO
            col['type'] = 'numericColumn'
        elif typ == 'cislo':
            col[':valueFormatter'] = _FMT_CISLO
            col['type'] = 'numericColumn'
        cols.append(col)
    return cols


def _podkladova_tabulka(pobocka: str, user_name: str) -> None:
    stav = {'radky': nacti_radky(pobocka), 'grids': {}, 'tab': None}

    @ui.refreshable
    def _grid_box():
        stav['grids'] = {}
        skupiny: dict[str, list[dict]] = {}
        for r in stav['radky']:
            skupiny.setdefault(_typ_bonusu(r), []).append(r)
        poradi = sorted(skupiny, key=lambda t: (t == _BEZ_TYPU, t.lower()))

        def _on_change(e):
            a = e.args or {}
            pole = a.get('colId')
            if pole not in POLE:
                return
            rid = (a.get('data') or {}).get('id')
            if not rid:
                return
            uloz_bunku(int(rid), pole, a.get('newValue'))
            for r in stav['radky']:
                if r['id'] == rid:
                    r[pole] = a.get('newValue')
                    break
            if pole == 'typ_bonusu':      # řádek patří do jiné podsekce
                _grid_box.refresh()

        def _panel(typ: str, radky: list[dict]):
            # grid se staví až pro aktivní záložku – jinak má nulovou šířku
            # a autosize sloupců na obsah se spočítá špatně
            grid = ui.aggrid({
                'columnDefs': _col_defs(),
                'rowData': radky,
                'defaultColDef': {'resizable': True, 'sortable': True, 'filter': True},
                'rowHeight': 32,
                'rowSelection': 'multiple',
                'suppressRowClickSelection': True,
                'suppressMovableColumns': True,
                'singleClickEdit': True,
                'stopEditingWhenCellsLoseFocus': True,
                ':onFirstDataRendered': _AUTOSIZE_FIT,
                ':onGridSizeChanged': _AUTOSIZE_FIT,
                ':getRowId': 'function(p){return String(p.data.id);}',
            }).classes('w-full').style(_GRID_STYLE)
            stav['grids'][typ] = grid
            grid.on('cellValueChanged', _on_change)

            async def _smaz(g=grid):
                vybrane = await g.get_selected_rows()
                ids = [int(r['id']) for r in vybrane if r.get('id')]
                if not ids:
                    ui.notify('Nejsou označené žádné řádky.', type='warning')
                    return
                pocet = smaz_radky(ids)
                stav['radky'] = nacti_radky(pobocka)
                _grid_box.refresh()
                ui.notify(f'Smazáno řádků: {pocet}', type='positive')
                intranet_logger.log_activity(
                    user_name, 'Bonusy AO',
                    f'{pobocka}: smazáno {pocet} řádků podkladové tabulky')

            with ui.row().classes('w-full justify-between items-center mt-2'):
                ui.label(f'Řádků: {len(radky)}').classes('text-sm text-gray-500')
                ui.button('Smazat označené', icon='delete', on_click=_smaz) \
                    .props('color=negative outline no-caps dense')

        if not poradi:
            ui.label('Žádné řádky.').classes('text-sm text-gray-500')
            return

        if stav.get('tab') not in poradi:
            stav['tab'] = poradi[0]

        def _prepni(e):
            if e.value and e.value != stav['tab']:
                stav['tab'] = e.value
                _grid_box.refresh()

        with ui.tabs(value=stav['tab'], on_change=_prepni) \
                .props('align=left active-color=primary indicator-color=primary') \
                .classes('w-full border-b border-gray-200'):
            for typ in poradi:
                ui.tab(typ, label=f'📋 {typ.upper()}')

        with ui.column().classes('w-full pt-2 gap-0'):
            _panel(stav['tab'], skupiny[stav['tab']])

    def _skoc_na_konec():
        typ = _BEZ_TYPU
        grid = stav['grids'].get(typ)
        radky = [r for r in stav['radky'] if _typ_bonusu(r) == typ]
        if grid is None or not radky:
            return
        idx = len(radky) - 1
        grid.run_grid_method('ensureIndexVisible', idx, 'bottom')
        grid.run_grid_method('startEditingCell', {'rowIndex': idx, 'colKey': POLE[0]})

    def _pridej():
        if not pridej_radek(pobocka):
            ui.notify('Řádek se nepodařilo přidat.', type='negative')
            return
        stav['radky'] = nacti_radky(pobocka)
        stav['tab'] = _BEZ_TYPU          # nový řádek je prázdný → záložka bez typu
        _grid_box.refresh()
        # grid se po refreshi staví znovu – skok až je na klientovi
        ui.timer(0.3, _skoc_na_konec, once=True)

    # ── dialog importu ──
    imp = {'raw': None, 'listy': [], 'nazev': ''}
    with ui.dialog() as dlg_import, ui.card().classes('w-[560px]'):
        ui.label(f'Import podkladové tabulky – {pobocka}').classes('text-lg font-bold')
        ui.label('Očekává se .xlsx se stejnými hlavičkami jako vzorová podkladová '
                 'tabulka (hlavičky v prvním řádku listu).') \
            .classes('text-xs text-gray-500')

        list_sel = ui.select([], label='List v sešitu').classes('w-full')
        nahradit = ui.checkbox('Před importem smazat stávající řádky pobočky')

        async def _on_upload(e):
            zdroj = next((getattr(e, a, None) for a in ('content', 'file', 'stream', 'data')
                          if hasattr(getattr(e, a, None), 'read')), None)
            if zdroj is None:
                ui.notify('Nepodařilo se načíst obsah souboru.', type='negative')
                return
            raw = zdroj.read()
            if inspect.isawaitable(raw):
                raw = await raw
            try:
                listy = await asyncio.to_thread(nacti_listy, raw)
            except Exception as exc:
                ui.notify(f'Soubor nejde otevřít: {exc}', type='negative')
                return
            imp['raw'] = raw
            imp['listy'] = listy
            list_sel.options = listy
            list_sel.value = listy[0] if listy else None
            list_sel.update()
            ui.notify(f'Načteno, listů: {len(listy)}', type='positive')

        ui.upload(on_upload=_on_upload, auto_upload=True, max_file_size=50_000_000,
                  label='Vybrat .xlsx soubor').props('accept=.xlsx').classes('w-full')

        async def _importuj():
            if not imp['raw'] or not list_sel.value:
                ui.notify('Nejdřív nahrajte soubor a vyberte list.', type='warning')
                return
            radky, err = await asyncio.to_thread(
                parsuj_excel, imp['raw'], list_sel.value, pobocka)
            if err:
                ui.notify(f'Chyba: {err}', type='negative', timeout=10000)
                return
            if not radky:
                ui.notify('List neobsahuje žádné datové řádky.', type='warning')
                return
            pocet = await asyncio.to_thread(uloz_import, pobocka, radky, nahradit.value)
            stav['radky'] = nacti_radky(pobocka)
            _grid_box.refresh()
            dlg_import.close()
            ui.notify(f'Importováno řádků: {pocet}', type='positive')
            intranet_logger.log_activity(
                user_name, 'Bonusy AO',
                f'{pobocka}: import podkladové tabulky – {pocet} řádků '
                f'(list {list_sel.value}{", nahrazeno" if nahradit.value else ""})')

        with ui.row().classes('w-full justify-end gap-2 mt-2'):
            ui.button('Zavřít', on_click=dlg_import.close).props('flat no-caps')
            ui.button('Importovat', icon='upload', on_click=_importuj).props('color=primary no-caps')

    with ui.column().classes('w-full gap-2'):
        with ui.row().classes('w-full justify-end gap-2'):
            ui.button('Přidat řádek', icon='add', on_click=_pridej) \
                .props('color=primary outline no-caps dense')
            ui.button('Import z Excelu', icon='upload_file', on_click=dlg_import.open) \
                .props('color=teal outline no-caps dense')
        _grid_box()


# ─── Data (SLK/XLSX z pobočkových složek) ─────────────────────────────────────

DATA_DIR = 'Bonusy'                 # kořen vedle aplikace: Bonusy/<pobočka>/Edit|Zpracovano
TABULKA_DATA = 'bonusy_ao_data'
TABULKA_KONTAKTY = 'bonusy_ao_kontakty'

# (db pole, hlavička, sql typ, šířka v gridu)
DATA_SLOUPCE = [
    ('ico',              'IČO',                       'VARCHAR(20)',   110),
    ('jmeno',            'Jméno',                     'VARCHAR(200)',  200),
    ('k_jmeno',          'K. Jméno',                  'VARCHAR(200)',  200),
    ('zakazka',          'Zakázka',                   'VARCHAR(30)',    90),
    ('typ_dokladu',      'Typ dokladu',               'VARCHAR(30)',   110),
    ('cislo_dokladu',    'Číslo dokladu',             'VARCHAR(30)',   120),
    ('datum',            'Datum zdanitelného plnění', 'DATE',          140),
    ('ext_id',           'Ext.ID',                    'VARCHAR(30)',   110),
    ('kod_zbozi',        'Kód zboží',                 'VARCHAR(30)',   100),
    ('nazev_zbozi',      'Název zboží',               'VARCHAR(150)',  220),
    ('skupina2',         'Skupina2',                  'VARCHAR(20)',    90),
    ('skupina2_popis',   'Skupina2-popis',            'VARCHAR(150)',  180),
    ('nadr_skup2',       'Nadř.skup.2',               'VARCHAR(20)',   100),
    ('nadr_skup2_popis', 'Nadř.skup.2 pop.',          'VARCHAR(150)',  180),
    ('dodavatel_popis',  'Dodavatel-popis',           'VARCHAR(150)',  180),
    ('sazba_dph',        'Sazba DPH',                 'VARCHAR(10)',    90),
    ('pouzita_ic',       'Použitá IC',                'VARCHAR(30)',   110),
    ('celk_bez_dph',     'Celk.bez DPH',              'DECIMAL(15,2)', 120),
]
DATA_POLE = [f for f, _h, _s, _w in DATA_SLOUPCE]
DATA_POLE_SQL = ', '.join(f'`{f}`' for f in DATA_POLE)

# pořadí sloupců ve zdrojovém souboru (hlavička je 5. řádek, data od 6.)
ZDROJ_POLE = [
    'ico', 'zakazka', 'typ_dokladu', 'cislo_dokladu', 'datum', 'ext_id', 'kod_zbozi',
    'nazev_zbozi', 'skupina2', 'skupina2_popis', 'nadr_skup2', 'nadr_skup2_popis',
    'dodavatel_popis', 'sazba_dph', 'pouzita_ic', 'celk_bez_dph',
]
PRVNI_DATOVY_RADEK = 6
VYNECHAT_DODAVATELE = {'fu', 'obaly'}       # DOD. popis „FÚ.“ a „Obaly“ → smazat řádky


def _slozka(pobocka: str, pod: str) -> str:
    cesta = os.path.join(DATA_DIR, pobocka, pod)
    os.makedirs(cesta, exist_ok=True)
    return cesta


def _datum_hodnota(hodnota):
    if isinstance(hodnota, datetime.datetime):
        return hodnota.date()
    if isinstance(hodnota, datetime.date):
        return hodnota
    txt = str(hodnota or '').strip().replace('/', '.')      # „/“ → „.“
    if not txt:
        return None
    for fmt in ('%d.%m.%y', '%d.%m.%Y', '%Y-%m-%d'):
        try:
            return datetime.datetime.strptime(txt, fmt).date()
        except ValueError:
            continue
    return None


def _cislo_hodnota(hodnota):
    if isinstance(hodnota, (int, float)):
        return float(hodnota)
    txt = str(hodnota or '').strip().replace(' ', '').replace(',', '.')
    try:
        return float(txt)
    except ValueError:
        return None


_SLK_BUNKA = re.compile(r'C;(?:Y(\d+);)?(?:X(\d+);)?K(.*)')


def _radky_slk(cesta: str, pokrok=None):
    y = x = 0
    radek: dict = {}
    velikost = max(os.path.getsize(cesta), 1)       # postup podle přečtených znaků
    nacteno = i_line = 0
    with io.open(cesta, 'r', encoding='cp1250', errors='replace') as f:
        for line in f:
            if pokrok:
                nacteno += len(line)
                i_line += 1
                if i_line % 2000 == 0:
                    pokrok(min(nacteno / velikost, 1.0))
            if not line.startswith('C;'):
                continue
            m = _SLK_BUNKA.match(line.rstrip('\r\n'))
            if not m:
                continue
            sy, sx, k = m.groups()
            if sy and int(sy) != y:
                if radek:
                    yield y, radek
                radek = {}
                y = int(sy)
            if sx:
                x = int(sx)
            v = k.strip()
            if v.startswith('"'):
                v = v[1:-1] if v.endswith('"') else v[1:]
            # SLK: středník v textu je escapovaný zdvojením
            radek[x] = v.replace(';;', ';').strip()
    if radek:
        yield y, radek


def _radky_xlsx(cesta: str, pokrok=None):
    import openpyxl
    wb = openpyxl.load_workbook(cesta, read_only=True, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        celkem = ws.max_row or 0
        for y, radek in enumerate(ws.iter_rows(values_only=True), 1):
            if pokrok and celkem and y % 500 == 0:
                pokrok(min(y / celkem, 1.0))
            yield y, {i: v for i, v in enumerate(radek, 1) if v is not None}
    finally:
        wb.close()


def zdrojove_radky(cesta: str, pokrok=None):
    """Datové řádky zdroje – prvních 5 řádků (4 hlavičkové + názvy sloupců) pryč.

    pokrok: volitelný callback(0.0–1.0) s podílem načteného zdroje.
    """
    fn = _radky_slk if cesta.lower().endswith('.slk') else _radky_xlsx
    for y, radek in fn(cesta, pokrok):
        if y >= PRVNI_DATOVY_RADEK:
            yield radek
    if pokrok:
        pokrok(1.0)


def uprav_radek(zdroj: dict, kontakty: dict) -> list | None:
    """Pravidla úpravy. None = řádek se zahazuje."""
    hod = {}
    for i, pole in enumerate(ZDROJ_POLE, 1):
        v = zdroj.get(i)
        hod[pole] = v.strip() if isinstance(v, str) else v
    ico = str(hod['ico'] or '').strip()
    if not ico or _norm(ico) == 'celkem':
        return None                                     # prázdný / součtový řádek
    if ico == '00000000':
        ico = str(hod['ext_id'] or '').strip()          # Ext.ID → IČO
    if _norm(hod['typ_dokladu']).replace('.', '').startswith('zavst'):
        return None                                     # ZAVStrvzenka
    if _norm(hod['dodavatel_popis']).replace('.', '') in VYNECHAT_DODAVATELE:
        return None                                     # FÚ. / Obaly
    jmeno, k_jmeno = kontakty.get(ico, ('', ''))        # SVYHLEDAT podle IČO
    hod['ico'] = ico
    hod['jmeno'] = jmeno
    hod['k_jmeno'] = k_jmeno
    hod['datum'] = _datum_hodnota(hod['datum'])
    hod['celk_bez_dph'] = _cislo_hodnota(hod['celk_bez_dph'])
    return [None if hod.get(p) == '' else hod.get(p) for p in DATA_POLE]


def _obdobi_z_nazvu(nazev: str) -> str:
    m = re.search(r'(\d{2})[_\-. ](\d{4})', nazev)
    if m:
        return f'{m.group(2)}-{m.group(1)}'
    return datetime.date.today().strftime('%Y-%m')


# ─── Kontaktní údaje VO (číselník pro Jméno / K. jméno) ──────────────────────

def parsuj_kontakty(raw: bytes, nazev: str) -> list[tuple]:
    if nazev.lower().endswith('.xlsb'):
        import pyxlsb
        radky = []
        with pyxlsb.open_workbook(io.BytesIO(raw)) as wb:
            with wb.get_sheet(1) as ws:
                for r in ws.rows():
                    radky.append([c.v for c in r])
    else:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        try:
            ws = wb[wb.sheetnames[0]]
            radky = [list(r) for r in ws.iter_rows(values_only=True)]
        finally:
            wb.close()

    def klic(v):
        return _norm(v).replace(' ', '').replace('.', '')

    mapa, start = {}, None
    for idx, r in enumerate(radky):
        volne = {klic(v): i for i, v in enumerate(r) if v not in (None, '')}
        if 'ico' in volne and 'jmeno' in volne:
            mapa, start = volne, idx + 1
            break
    if start is None:
        raise ValueError('V souboru nejsou hlavičky IČO / Jméno.')
    vysledek, videna = [], set()
    for r in radky[start:]:
        ico = str(r[mapa['ico']] or '').strip() if mapa['ico'] < len(r) else ''
        if not ico:
            continue
        def bunka(k):
            i = mapa.get(k)
            return str(r[i] or '').strip()[:200] if i is not None and i < len(r) else ''
        ico = ico[:20]
        if ico in videna:
            continue            # SVYHLEDAT bere první výskyt IČO
        videna.add(ico)
        vysledek.append((ico, bunka('jmeno'), bunka('kjmeno')))
    return vysledek


def uloz_kontakty(dvojice: list[tuple]) -> int:
    conn = intranet_data.get_db_connection()
    if not conn:
        return 0
    try:
        cur = conn.cursor()
        cur.execute(f'DELETE FROM {TABULKA_KONTAKTY}')
        sql = f'INSERT INTO {TABULKA_KONTAKTY} (ico, jmeno, k_jmeno) VALUES (%s, %s, %s)'
        for i in range(0, len(dvojice), 5000):
            cur.executemany(sql, dvojice[i:i + 5000])
        conn.commit()
        cur.close()
        return len(dvojice)
    finally:
        conn.close()


def nacti_kontakty() -> dict:
    conn = intranet_data.get_db_connection()
    if not conn:
        return {}
    try:
        cur = conn.cursor()
        cur.execute(f'SELECT ico, jmeno, k_jmeno FROM {TABULKA_KONTAKTY}')
        return {r[0]: (r[1] or '', r[2] or '') for r in cur.fetchall()}
    finally:
        conn.close()


def pocet_kontaktu() -> int:
    conn = intranet_data.get_db_connection()
    if not conn:
        return 0
    try:
        cur = conn.cursor()
        cur.execute(f'SELECT COUNT(*) FROM {TABULKA_KONTAKTY}')
        return cur.fetchone()[0]
    finally:
        conn.close()


# ─── Zpracování souboru ──────────────────────────────────────────────────────

def inicializace_bonusy_data_db():
    conn = intranet_data.get_db_connection()
    if not conn:
        return
    try:
        sloupce_ddl = ',\n                '.join(
            f'`{f}` {t}' for f, _h, t, _w in DATA_SLOUPCE)
        cur = conn.cursor()
        cur.execute(f"""
            CREATE TABLE IF NOT EXISTS {TABULKA_DATA} (
                id      BIGINT AUTO_INCREMENT PRIMARY KEY,
                pobocka_klic VARCHAR(50) NOT NULL,
                obdobi  VARCHAR(7) NOT NULL,
                {sloupce_ddl},
                INDEX idx_pob_obd (pobocka_klic, obdobi),
                INDEX idx_ico (ico)
            ) CHARACTER SET utf8mb4 COLLATE utf8mb4_unicode_ci
        """)
        cur.execute(f"""
            CREATE TABLE IF NOT EXISTS {TABULKA_KONTAKTY} (
                ico     VARCHAR(20) COLLATE utf8mb4_bin NOT NULL PRIMARY KEY,
                jmeno   VARCHAR(200),
                k_jmeno VARCHAR(200)
            ) CHARACTER SET utf8mb4 COLLATE utf8mb4_unicode_ci
        """)
        conn.commit()
        cur.close()
    except Exception as e:
        print(f'[bonusy_ao] Chyba inicializace DB dat: {e}')
    finally:
        conn.close()


def zpracuj_soubor(pobocka: str, nazev: str, pokrok=None) -> tuple[int, str, str]:
    """Načte soubor z Edit, aplikuje pravidla, uloží do DB a do Zpracovano.

    Období řádku = měsíc data zdanitelného plnění; bez data → z názvu souboru.
    Vrací (pocet, cil, prevazujici_obdobi).
    """
    import openpyxl
    from openpyxl.cell import WriteOnlyCell
    zdroj = os.path.join(_slozka(pobocka, 'Edit'), nazev)
    zaloha_obdobi = _obdobi_z_nazvu(nazev)
    pocty_obdobi: dict[str, int] = {}
    smazana: set[str] = set()
    kontakty = nacti_kontakty()
    conn = intranet_data.get_db_connection()
    if not conn:
        raise RuntimeError('Databáze není dostupná.')
    wb = openpyxl.Workbook(write_only=True)
    ws = wb.create_sheet(title='data')
    ws.append([h for _f, h, _s, _w in DATA_SLOUPCE])
    sql = (f'INSERT INTO {TABULKA_DATA} (pobocka_klic, obdobi, {DATA_POLE_SQL}) '
           f'VALUES ({", ".join(["%s"] * (len(DATA_POLE) + 2))})')
    davka, pocet = [], 0
    i_datum = DATA_POLE.index('datum')
    try:
        cur = conn.cursor()
        for zdrojovy in zdrojove_radky(zdroj, pokrok):
            radek = uprav_radek(zdrojovy, kontakty)
            if radek is None:
                continue
            bunky = list(radek)
            datum = bunky[i_datum]
            if isinstance(datum, (datetime.date, datetime.datetime)):
                obdobi = datum.strftime('%Y-%m')
                b = WriteOnlyCell(ws, value=datum)
                b.number_format = 'DD.MM.YYYY'
                bunky[i_datum] = b
            else:
                obdobi = zaloha_obdobi
            if obdobi not in smazana:      # staré řádky období pryč před importem
                cur.executemany(sql, davka)
                conn.commit()
                davka = []
                cur.execute(
                    f'DELETE FROM {TABULKA_DATA} WHERE pobocka_klic=%s AND obdobi=%s',
                    (pobocka, obdobi))
                conn.commit()
                smazana.add(obdobi)
            pocty_obdobi[obdobi] = pocty_obdobi.get(obdobi, 0) + 1
            ws.append(bunky)
            davka.append((pobocka, obdobi, *radek))
            pocet += 1
            if len(davka) >= 5000:
                cur.executemany(sql, davka)
                conn.commit()
                davka = []
        if davka:
            cur.executemany(sql, davka)
            conn.commit()
        cur.close()
    finally:
        conn.close()
    cil = os.path.join(_slozka(pobocka, 'Zpracovano'),
                       f'{os.path.splitext(nazev)[0]}_zpracovano.xlsx')
    wb.save(cil)
    if zdroj.lower().endswith('.slk'):   # zdrojový SLK po zpracování pryč
        os.remove(zdroj)
    hlavni = max(pocty_obdobi, key=pocty_obdobi.get) if pocty_obdobi else zaloha_obdobi
    return pocet, cil, hlavni


def nacti_obdobi(pobocka: str) -> list[str]:
    """Období pobočky v DB, nejnovější první. DISTINCT přes prefix indexu –
    COUNT(*) by skenoval miliony řádků a zablokoval UI."""
    conn = intranet_data.get_db_connection()
    if not conn:
        return []
    try:
        cur = conn.cursor()
        cur.execute(
            f'SELECT DISTINCT pobocka_klic, obdobi FROM {TABULKA_DATA} '
            f'WHERE pobocka_klic=%s', (pobocka,))
        out = sorted((r[1] for r in cur.fetchall()), reverse=True)
        cur.close()
        return out
    finally:
        conn.close()


def smaz_obdobi(pobocka: str, obdobi: str) -> int:
    conn = intranet_data.get_db_connection()
    if not conn:
        return 0
    try:
        cur = conn.cursor()
        cur.execute(f'DELETE FROM {TABULKA_DATA} WHERE pobocka_klic=%s AND obdobi=%s',
                    (pobocka, obdobi))
        pocet = cur.rowcount
        conn.commit()
        cur.close()
        return pocet
    finally:
        conn.close()


def nacti_data_nahled(pobocka: str, obdobi: str | None = None,
                      limit: int = 500) -> tuple[list[dict], int]:
    conn = intranet_data.get_db_connection()
    if not conn:
        return [], 0
    kde = 'pobocka_klic=%s' + (' AND obdobi=%s' if obdobi else '')
    par = (pobocka, obdobi) if obdobi else (pobocka,)
    try:
        cur = conn.cursor(dictionary=True)
        cur.execute(f'SELECT COUNT(*) AS c FROM {TABULKA_DATA} WHERE {kde}', par)
        celkem = cur.fetchone()['c']
        cur.execute(
            f'SELECT {DATA_POLE_SQL} FROM {TABULKA_DATA} WHERE {kde} '
            f'ORDER BY id LIMIT {int(limit)}', par)
        radky = cur.fetchall()
        cur.close()
        for r in radky:
            if isinstance(r.get('datum'), (datetime.date, datetime.datetime)):
                r['datum'] = r['datum'].strftime('%d.%m.%Y')
            if r.get('celk_bez_dph') is not None:
                r['celk_bez_dph'] = float(r['celk_bez_dph'])
        return radky, celkem
    finally:
        conn.close()


# ─── Výpočet bonusů (port logiky z App.java) ──────────────────────────────────

# klíč, popisek, předpona typu bonusu v podkladu (normalizovaná: „fakturačně“ i
# „fakturační“). Ostatní typy (zboží, odd + zápočet…) se nepočítají – jako v Javě.
BONUS_TYPY = [
    ('fakturacni', 'Fakturační', 'fakturac'),
    ('konto_o',    'Konto O',    'konto o'),
    ('konto_r',    'Konto R',    'konto r'),
]


def _bonus_typ_klic(hodnota) -> str | None:
    t = _norm(hodnota)
    for klic, _lbl, predpona in BONUS_TYPY:
        if t.startswith(predpona):
            return klic
    return None

# výstupní sloupce = data + 4 dopočtené
BONUS_SLOUPCE = DATA_SLOUPCE + [
    ('bonus_proc', 'bonus v %',           'procento',      100),
    ('bonus_kc',   'bonus v Kč bez DPH',  'DECIMAL(15,2)', 150),
    ('skupina',    'skupina',             'text',          120),
    ('min_odber',  'minimální odběr',     'DECIMAL(15,2)', 130),
]

# indexy v datovém řádku (pořadí DATA_POLE)
_I_ICO, _I_TYP_DOKL, _I_DATUM = 0, 4, 6
_I_NADR_POPIS, _I_DOD_POPIS, _I_POUZITA_IC, _I_CASTKA = 13, 14, 16, 17

_BONUS_CACHE: dict[str, dict] = {}      # náhled posledního výpočtu (per pobočka|okno|režim)

_MAX_RADKU_LIST = 1_000_000             # limit XLSX je 1 048 576 vč. hlavičky

# režimy okna výpočtu (klíč, popisek, délka v měsících)
BONUS_REZIMY = [
    ('mesic',    'Měsíc',        1),
    ('kvartal',  'Kvartál',      3),
    ('pololeti', 'Pololetí',     6),
    ('rok',      'Rok',         12),
    ('podklad',  'Dle podkladu', 0),
]
_REZIM_DELKA = {k: n for k, _l, n in BONUS_REZIMY if n}
_REZIM_LBL = {k: lbl for k, lbl, _n in BONUS_REZIMY}
# sloupec Interval v podkladu -> režim okna
_IVL_REZIM = {'mesicni': 'mesic', 'kvartalni': 'kvartal',
              'pololetni': 'pololeti', 'rocni': 'rok'}


def _ivl_rezim(hodnota) -> str:
    return _IVL_REZIM.get(_norm(hodnota), 'mesic')


def _okno_mesice(obdobi: str, rezim: str) -> tuple[str, list[str]]:
    """Kalendářní okno kolem období 'YYYY-MM'. Vrací (popisek, [měsíce])."""
    rok, mes = int(obdobi[:4]), int(obdobi[5:7])
    delka = _REZIM_DELKA.get(rezim, 1)
    if delka == 1:
        return obdobi, [obdobi]
    zac = ((mes - 1) // delka) * delka + 1
    mesice = [f'{rok}-{m:02d}' for m in range(zac, zac + delka)]
    if delka == 3:
        popis = f'{rok}-Q{(zac - 1) // 3 + 1}'
    elif delka == 6:
        popis = f'{rok}-H{(zac - 1) // 6 + 1}'
    else:
        popis = str(rok)
    return popis, mesice


def _bonus_datum(txt) -> datetime.date | None:
    txt = str(txt or '').strip()
    if not txt:
        return None
    try:
        return datetime.datetime.strptime(txt, '%d.%m.%Y').date()
    except ValueError:
        return None


def _bonus_mapy(pobocka: str) -> dict[str, dict]:
    """{typ bonusu: {IČO: [řádky podkladu]}} – jedno IČO může mít víc období
    platnosti (Od–Do). Seřazeno od nejnovějšího Od, aby při překryvu vyhrál
    novější řádek. Do řádku dopočte _od, _do (platnost) a _rezim (z Intervalu)."""
    mapy: dict[str, dict] = {k: {} for k, _l, _n in BONUS_TYPY}
    for r in nacti_radky(pobocka):
        klic = _bonus_typ_klic(r.get('typ_bonusu'))
        ico = str(r.get('ico_zadane') or '').strip()
        if klic and ico:
            r['_od'] = _bonus_datum(r.get('datum_od'))
            r['_do'] = _bonus_datum(r.get('datum_do'))
            r['_rezim'] = _ivl_rezim(r.get('interval'))
            mapy[klic].setdefault(ico, []).append(r)
    for mapa in mapy.values():
        for radky in mapa.values():
            radky.sort(key=lambda p: p['_od'] or datetime.date.min, reverse=True)
    return mapy


def _mapy_pro_rezim(mapy: dict[str, dict], rezim: str) -> dict[str, dict]:
    """Pevný režim počítá jen řádky podkladu s odpovídajícím Intervalem.
    Režim 'podklad' bere vše, každý řádek ve svém okně."""
    if rezim == 'podklad':
        return mapy
    out: dict[str, dict] = {}
    for k, mapa in mapy.items():
        out[k] = {i: [p for p in radky if p['_rezim'] == rezim]
                  for i, radky in mapa.items()}
        out[k] = {i: radky for i, radky in out[k].items() if radky}
    return out


def _podklad_pro(mapa: dict, ico: str, datum):
    """Řádek podkladu IČO platný k datu zdanitelného plnění; jinak None."""
    for p in mapa.get(ico, ()):
        if not _mimo_platnost(p, datum):
            return p
    return None


def _mimo_platnost(p: dict, datum) -> bool:
    """Datum zdanitelného plnění mimo Od–Do podkladu (obě hranice včetně)."""
    if isinstance(datum, datetime.datetime):
        datum = datum.date()
    if p['_od'] and datum < p['_od']:
        return True
    if p['_do'] and datum > p['_do']:
        return True
    return False


def _okno_pro_vypocet(pobocka: str, obdobi: str, rezim: str) -> tuple[str, list[str]]:
    """Okno pevného režimu; u 'podklad' sjednocení oken všech intervalů v podkladu."""
    if rezim != 'podklad':
        return _okno_mesice(obdobi, rezim)
    mesice: set[str] = set()
    for r in nacti_radky(pobocka):
        if _bonus_typ_klic(r.get('typ_bonusu')):
            mesice.update(_okno_mesice(obdobi, _ivl_rezim(r.get('interval')))[1])
    return f'{obdobi}-podklad', sorted(mesice or [obdobi])


def bonus_predkontrola(pobocka: str, obdobi: str, rezim: str) -> dict:
    """Kontroly před výpočtem: uzavřenost okna, data bez datumu, podklad bez platnosti."""
    popis, mesice = _okno_pro_vypocet(pobocka, obdobi, rezim)
    v_db = set(nacti_obdobi(pobocka))
    chybi = [m for m in mesice if m not in v_db]
    bez_data = 0
    if not chybi:
        conn = intranet_data.get_db_connection()
        if not conn:
            raise RuntimeError('Databáze není dostupná.')
        try:
            cur = conn.cursor()
            cur.execute(
                f'SELECT COUNT(*) FROM {TABULKA_DATA} WHERE pobocka_klic=%s '
                f'AND obdobi IN ({", ".join(["%s"] * len(mesice))}) AND datum IS NULL',
                (pobocka, *mesice))
            bez_data = cur.fetchone()[0] or 0
            cur.close()
        finally:
            conn.close()
    mapy = _mapy_pro_rezim(_bonus_mapy(pobocka), rezim)
    podklad = [p for mapa in mapy.values() for radky in mapa.values() for p in radky]
    bez_platnosti = sum(1 for p in podklad if not p['_od'] and not p['_do'])
    return {'popis': popis, 'mesice': mesice, 'chybi_obdobi': chybi,
            'bez_data': bez_data, 'bez_platnosti': bez_platnosti,
            'podklad': len(podklad)}


def _bonus_vyhodit(p: dict, d: tuple) -> bool:
    """A/N příznaky podkladu, které řádek z výstupu úplně odstraní."""
    typ_dokl = str(d[_I_TYP_DOKL] or '').strip()
    nadr = str(d[_I_NADR_POPIS] or '')
    dod = str(d[_I_DOD_POPIS] or '')
    ic = str(d[_I_POUZITA_IC] or '')
    return bool(
        (p['obrat_dp'] == 'N' and typ_dokl == 'Dr.prodej')
        or (p['obrat_akce'] == 'N' and 'PC' in ic)
        or (p['tabak'] == 'N' and 'tabák' in nadr)
        or (p['pivo'] == 'N' and 'pivo' in nadr)
        or (p['ceniny'] == 'N' and 'provize' in nadr)
        or (p['nektar_natura'] == 'N' and 'Nektar' in dod)
        or (p['eso'] == 'N' and dod.startswith('ESO'))
        or (p['premier_wines'] == 'N' and 'Premier Wines' in dod)
    )


def _bonus_nulovat(p: dict, d: tuple) -> bool:
    """Řádek zůstává, ale bonus je nulový."""
    return bool(
        (p['bonus_dp'] == 'N' and str(d[_I_TYP_DOKL] or '').strip() == 'Dr.prodej')
        or (p['bonus_akce'] == 'N' and 'PC' in str(d[_I_POUZITA_IC] or ''))
    )


def _bonus_data_radky(conn, pobocka: str, mesice: list[str]):
    """Streamuje datové řádky okna po dávkách (ať se 150k řádků nevejde naráz)."""
    cur = conn.cursor()
    cur.execute(f'SELECT {DATA_POLE_SQL} FROM {TABULKA_DATA} '
                f'WHERE pobocka_klic=%s AND obdobi IN '
                f'({", ".join(["%s"] * len(mesice))}) ORDER BY obdobi, id',
                (pobocka, *mesice))
    try:
        while True:
            davka = cur.fetchmany(5000)
            if not davka:
                return
            yield from davka
    finally:
        cur.close()


def _bonus_nazev(klic: str, popis: str) -> str:
    return f'bonusy_{klic}_{popis}.xlsx'


def _bonus_mesic(datum) -> str:
    return datum.strftime('%Y-%m')


def _novy_list(wb, hlavicka: list):
    """Nový list s hlavičkou – kvůli limitu 1 048 576 řádků na list."""
    poradi = len(wb.worksheets) + 1
    w = wb.create_sheet(title='bonusy' if poradi == 1 else f'bonusy_{poradi}')
    w.append(hlavicka)
    return w


def spocitej_bonusy(pobocka: str, obdobi: str, rezim: str = 'mesic', pokrok=None) -> dict:
    """Dopočte bonusy pro 3 typy za zvolené okno, zapíše XLSX do Zpracovano.
    Vrací {'pocty': {...}, 'radky': {...náhled}, 'soubory': {...}}."""
    import openpyxl
    from openpyxl.cell import WriteOnlyCell

    mapy = _mapy_pro_rezim(_bonus_mapy(pobocka), rezim)
    popis, mesice = _okno_pro_vypocet(pobocka, obdobi, rezim)
    # okno pro součet min. odběru: pevné pro všechny, nebo per řádek podkladu
    okna: dict[int, set[str]] = {}
    okno_pevne = set(mesice)

    def _okno_radku(p: dict) -> set[str]:
        if rezim != 'podklad':
            return okno_pevne
        ms = okna.get(id(p))
        if ms is None:
            ms = set(_okno_mesice(obdobi, p['_rezim'])[1])
            okna[id(p)] = ms
        return ms

    hotovo = _slozka(pobocka, 'Zpracovano')
    conn = intranet_data.get_db_connection()
    if not conn:
        raise RuntimeError('Databáze není dostupná.')
    try:
        cur = conn.cursor()
        cur.execute(f'SELECT COUNT(*) FROM {TABULKA_DATA} WHERE pobocka_klic=%s '
                    f'AND obdobi IN ({", ".join(["%s"] * len(mesice))})',
                    (pobocka, *mesice))
        celkem = cur.fetchone()[0] or 0
        cur.close()

        # 1. průchod – obrat za klíč (skupina, jinak IČO) a měsíc, pro min. odběr
        soucty = {k: {} for k, _l, _n in BONUS_TYPY}
        for i, d in enumerate(_bonus_data_radky(conn, pobocka, mesice)):
            ico = str(d[_I_ICO] or '').strip()
            mes = _bonus_mesic(d[_I_DATUM])
            for klic, mapa in mapy.items():
                p = _podklad_pro(mapa, ico, d[_I_DATUM])
                if p is None or _bonus_vyhodit(p, d):
                    continue
                if mes not in _okno_radku(p):
                    continue
                kl = str(p['skupina'] or '').strip() or ico
                po_mesicich = soucty[klic].setdefault(kl, {})
                po_mesicich[mes] = po_mesicich.get(mes, 0.0) + float(d[_I_CASTKA] or 0)
            if pokrok and celkem and i % 2000 == 0:
                pokrok(i / (2 * celkem))

        # 2. průchod – výstupní řádky + zápis
        hlavicka = [h for _f, h, _t, _w in BONUS_SLOUPCE]
        wb = {k: openpyxl.Workbook(write_only=True) for k, _l, _n in BONUS_TYPY}
        ws = {k: _novy_list(wb[k], hlavicka) for k, _l, _n in BONUS_TYPY}
        na_listu = {k: 0 for k, _l, _n in BONUS_TYPY}
        pocty = {k: 0 for k, _l, _n in BONUS_TYPY}
        nahled = {k: [] for k, _l, _n in BONUS_TYPY}

        for i, d in enumerate(_bonus_data_radky(conn, pobocka, mesice)):
            ico = str(d[_I_ICO] or '').strip()
            mes = _bonus_mesic(d[_I_DATUM])
            for klic, mapa in mapy.items():
                p = _podklad_pro(mapa, ico, d[_I_DATUM])
                if p is None or _bonus_vyhodit(p, d):
                    continue
                okno = _okno_radku(p)
                if mes not in okno:
                    continue
                castka = float(d[_I_CASTKA] or 0)
                proc = p['procento_bonusu'] or 0.0
                if _bonus_nulovat(p, d):
                    proc = 0.0
                kc = round(castka * proc, 2)
                skupina = str(p['skupina'] or '').strip()
                min_odber = p['minimalni_odber']
                if min_odber is not None:
                    po_mesicich = soucty[klic].get(skupina or ico, {})
                    if sum(po_mesicich.get(m, 0.0) for m in okno) < min_odber:
                        kc = 0.0
                radek = list(d) + [proc, kc, skupina, min_odber]

                if na_listu[klic] >= _MAX_RADKU_LIST:   # limit XLSX -> další list
                    ws[klic] = _novy_list(wb[klic], hlavicka)
                    na_listu[klic] = 0
                bunky = list(radek)
                datum = bunky[_I_DATUM]
                if isinstance(datum, (datetime.date, datetime.datetime)):
                    b = WriteOnlyCell(ws[klic], value=datum)
                    b.number_format = 'DD.MM.YYYY'
                    bunky[_I_DATUM] = b
                b = WriteOnlyCell(ws[klic], value=proc)
                b.number_format = '0.00%'
                bunky[len(DATA_POLE)] = b
                ws[klic].append(bunky)
                na_listu[klic] += 1
                pocty[klic] += 1
                if len(nahled[klic]) < 500:
                    n = dict(zip(DATA_POLE, d))
                    if isinstance(n.get('datum'), (datetime.date, datetime.datetime)):
                        n['datum'] = n['datum'].strftime('%d.%m.%Y')
                    n['celk_bez_dph'] = castka
                    n.update(bonus_proc=proc, bonus_kc=kc,
                             skupina=skupina, min_odber=min_odber)
                    nahled[klic].append(n)
            if pokrok and celkem and i % 2000 == 0:
                pokrok(0.5 + i / (2 * celkem))
    finally:
        conn.close()

    soubory = {}
    for klic, _lbl, _n in BONUS_TYPY:
        cil = os.path.join(hotovo, _bonus_nazev(klic, popis))
        wb[klic].save(cil)
        soubory[klic] = cil
    if pokrok:
        pokrok(1.0)
    return {'pocty': pocty, 'radky': nahled, 'soubory': soubory,
            'popis': popis, 'mesice': mesice, 'rezim': rezim,
            'cas': datetime.datetime.now()}


def _bonus_col_defs() -> list[dict]:
    cols = []
    for f, h, typ, sirka in BONUS_SLOUPCE:
        col = {'headerName': h, 'field': f, 'width': sirka,
               'sortable': True, 'filter': True, 'headerTooltip': h}
        if typ.startswith('DECIMAL'):
            col['type'] = 'numericColumn'
            col['valueFormatter'] = _FMT_CISLO
        elif typ == 'procento':
            col['type'] = 'numericColumn'
            col['valueFormatter'] = _FMT_PROCENTO
        cols.append(col)
    return cols


@refreshable_na_klienta
def _bonusy_sekce(pobocka: str, user_name: str) -> None:
    inicializace_bonusy_data_db()
    hotovo = _slozka(pobocka, 'Zpracovano')
    obdobi_klice = nacti_obdobi(pobocka)
    stav_klic = f'bonusy_ao_bonus_obdobi_{pobocka}'
    obdobi_sel = app.storage.user.get(stav_klic)
    if obdobi_sel not in obdobi_klice:
        obdobi_sel = obdobi_klice[0] if obdobi_klice else None
        app.storage.user[stav_klic] = obdobi_sel
    rezim_klic = f'bonusy_ao_bonus_rezim_{pobocka}'
    rezim_sel = app.storage.user.get(rezim_klic)
    if rezim_sel not in [k for k, _l, _n in BONUS_REZIMY]:
        rezim_sel = 'mesic'
        app.storage.user[rezim_klic] = rezim_sel
    okno_popis, okno_mesice = (
        _okno_pro_vypocet(pobocka, obdobi_sel, rezim_sel) if obdobi_sel else ('', []))
    vysledek = _BONUS_CACHE.get(f'{pobocka}|{okno_popis}|{rezim_sel}')

    async def _spocitej():
        if not obdobi_sel:
            return
        try:
            kontrola = await asyncio.to_thread(
                bonus_predkontrola, pobocka, obdobi_sel, rezim_sel)
        except Exception as exc:
            ui.notify(f'Chyba kontroly: {exc}', type='negative', timeout=10000)
            return
        if kontrola['chybi_obdobi']:
            ui.notify(f'Okno {kontrola["popis"]} není uzavřené – nejdřív zpracujte '
                      f'období: {", ".join(kontrola["chybi_obdobi"])}',
                      type='negative', timeout=15000)
            return
        if not kontrola['podklad']:
            ui.notify(f'Podkladová tabulka nemá žádný řádek s intervalem '
                      f'{_REZIM_LBL.get(rezim_sel, rezim_sel)}.',
                      type='negative', timeout=10000)
            return
        if kontrola['bez_data']:
            ui.notify(f'{kontrola["bez_data"]} řádků dat nemá Datum zdanitelného '
                      f'plnění – nelze ověřit platnost bonusu, výpočet zastaven.',
                      type='negative', timeout=15000)
            return
        if kontrola['bez_platnosti']:
            with ui.dialog() as dlg_upoz, ui.card().classes('w-[460px]'):
                ui.label('Podklad bez platnosti').classes('text-lg font-bold')
                ui.label(f'{kontrola["bez_platnosti"]} řádků podkladové tabulky nemá '
                         f'vyplněno Od ani Do. Počítají se bez omezení platnosti.') \
                    .classes('text-sm text-gray-600')
                with ui.row().classes('w-full justify-end gap-2'):
                    ui.button('Zrušit', on_click=lambda: dlg_upoz.submit(False)) \
                        .props('flat no-caps')
                    ui.button('Pokračovat', on_click=lambda: dlg_upoz.submit(True)) \
                        .props('color=primary no-caps')
            if not await dlg_upoz:
                return
        stav_pokrok = {'f': 0.0}
        dlg, kruh, popisek = _prekryv_kolecko(f'Počítám bonusy {okno_popis}…')

        def _tik():
            proc = stav_pokrok['f'] * 100
            kruh.set_value(proc)
            popisek.set_text(f'{proc:.0f} %')

        dlg.open()
        casovac = ui.timer(0.2, _tik)
        try:
            out = await asyncio.to_thread(
                spocitej_bonusy, pobocka, obdobi_sel, rezim_sel,
                lambda f: stav_pokrok.update(f=f))
        except Exception as exc:
            ui.notify(f'Chyba výpočtu: {exc}', type='negative', timeout=10000)
            return
        finally:
            casovac.cancel()
            dlg.close()
        _BONUS_CACHE[f'{pobocka}|{out["popis"]}|{rezim_sel}'] = out
        ui.notify(f'Bonusy {out["popis"]} spočteny: ' + ', '.join(
            f'{lbl} {out["pocty"][k]}' for k, lbl, _n in BONUS_TYPY), type='positive')
        intranet_logger.log_activity(
            user_name, 'Bonusy AO',
            f'{pobocka}: spočteny bonusy za {out["popis"]} '
            f'({sum(out["pocty"].values())} řádků)')
        _bonusy_sekce.refresh()

    def _zmen_obdobi(e):
        app.storage.user[stav_klic] = e.value
        _bonusy_sekce.refresh()

    def _zmen_rezim(e):
        app.storage.user[rezim_klic] = e.value
        _bonusy_sekce.refresh()

    with ui.column().classes('w-full gap-2'):
        with ui.row().classes('w-full items-center gap-3'):
            ui.label('Bonusy').classes('text-lg font-bold text-gray-800')
            if obdobi_klice:
                ui.select(obdobi_klice,
                          value=obdobi_sel, label='Období', on_change=_zmen_obdobi) \
                    .props('dense outlined options-dense').style('min-width: 220px')
                ui.select({k: lbl for k, lbl, _n in BONUS_REZIMY},
                          value=rezim_sel, label='Interval', on_change=_zmen_rezim) \
                    .props('dense outlined options-dense').style('min-width: 160px')
                ui.button('Spočítat', icon='calculate', on_click=_spocitej) \
                    .props('color=primary outline no-caps dense')
            else:
                ui.label('Nejdřív zpracujte data v sekci Data.') \
                    .classes('text-sm text-gray-400')
            ui.space()
            for klic, lbl, _n in BONUS_TYPY:
                plna = os.path.join(hotovo, _bonus_nazev(klic, okno_popis))
                if obdobi_sel and os.path.isfile(plna):
                    ui.button(f'Stáhnout {lbl}', icon='download',
                              on_click=lambda _e, c=plna:
                                  ui.download.file(c, os.path.basename(c))) \
                        .props('color=green outline no-caps dense')
        if obdobi_sel:
            if rezim_sel == 'podklad':
                info = (f'Okno {okno_mesice[0]} – {okno_mesice[-1]} • délka per řádek '
                        f'dle sloupce Interval v podkladu')
            else:
                pocet_podkladu = sum(
                    len(radky) for mapa in
                    _mapy_pro_rezim(_bonus_mapy(pobocka), rezim_sel).values()
                    for radky in mapa.values())
                info = (f'Okno {okno_popis} ({okno_mesice[0]} – {okno_mesice[-1]}) • '
                        f'podklad: {pocet_podkladu} řádků s intervalem '
                        f'{_REZIM_LBL[rezim_sel]}')
            ui.label(info).classes('text-xs text-gray-500')
        if not vysledek:
            ui.label('Zatím nespočítáno – klikněte na Spočítat.') \
                .classes('text-sm text-gray-400')
            return
        ui.label(f'Spočteno {vysledek["cas"]:%d.%m.%Y %H:%M} • '
                 f'zdroj: podkladová tabulka + data {vysledek["popis"]}') \
            .classes('text-xs text-gray-500')
        with ui.tabs().props(
            'align=left active-color=primary indicator-color=primary'
        ).classes('w-full border-b border-gray-200') as taby:
            for klic, lbl, _n in BONUS_TYPY:
                ui.tab(klic, label=f'{lbl} ({vysledek["pocty"][klic]})')
        with ui.tab_panels(taby, value=BONUS_TYPY[0][0]).classes('w-full pt-2'):
            for klic, _lbl, _n in BONUS_TYPY:
                with ui.tab_panel(klic).classes('p-0'):
                    ui.aggrid({
                        'columnDefs': _bonus_col_defs(),
                        'rowData': vysledek['radky'][klic],
                        'defaultColDef': {'resizable': True, 'sortable': True,
                                          'filter': True},
                        'rowHeight': 32,
                        'suppressMovableColumns': True,
                        ':onFirstDataRendered': _AUTOSIZE_FIT,
                        ':onGridSizeChanged': _AUTOSIZE_FIT,
                    }).classes('w-full').style(_GRID_STYLE)
                    if vysledek['pocty'][klic] > len(vysledek['radky'][klic]):
                        ui.label(f'Náhled prvních {len(vysledek["radky"][klic])} '
                                 f'z {vysledek["pocty"][klic]} řádků – '
                                 f'celý výstup je v XLSX.') \
                            .classes('text-xs text-gray-500')


def _soubory(cesta: str) -> list[tuple]:
    out = []
    for nazev in sorted(os.listdir(cesta)):
        plna = os.path.join(cesta, nazev)
        if os.path.isfile(plna) and not nazev.startswith('~$'):
            out.append((nazev, os.path.getsize(plna)))
    return out


def _velikost(b: int) -> str:
    for jed in ('B', 'kB', 'MB', 'GB'):
        if b < 1024 or jed == 'GB':
            return f'{b:.0f} {jed}' if jed == 'B' else f'{b:.1f} {jed}'
        b /= 1024
    return f'{b:.1f} GB'


def _prekryv_kolecko(popis: str, procenta: bool = True):
    """Šedý překryv s kolečkem uprostřed. Vrací (dialog, kruh, popisek).
    procenta=False → kolečko se točí (délka operace není známá)."""
    with ui.dialog().props('persistent') as dlg, \
            ui.card().classes('bg-transparent shadow-none items-center gap-2'):
        kruh = ui.circular_progress(value=0, max=100, size='98px', show_value=False) \
            .props('thickness=0.2 color=primary track-color=grey-5')
        if not procenta:
            kruh.props('indeterminate')
        with kruh:
            popisek = ui.label('0 %' if procenta else '') \
                .classes('absolute-center text-lg font-bold')
        ui.label(popis).classes('text-white').style('font-size: 1.25rem')
    return dlg, kruh, popisek


def _data_col_defs() -> list[dict]:
    cols = []
    for f, h, typ, sirka in DATA_SLOUPCE:
        col = {'headerName': h, 'field': f, 'width': sirka,
               'sortable': True, 'filter': True, 'headerTooltip': h}
        if typ.startswith('DECIMAL'):
            col['type'] = 'numericColumn'
            col['valueFormatter'] = _FMT_CISLO
        cols.append(col)
    return cols


@refreshable_na_klienta
def _data_sekce(pobocka: str, user_name: str) -> None:
    inicializace_bonusy_data_db()
    edit = _slozka(pobocka, 'Edit')
    hotovo = _slozka(pobocka, 'Zpracovano')
    obdobi_klic = f'bonusy_ao_obdobi_{pobocka}'
    obdobi_klice = nacti_obdobi(pobocka)
    obdobi_sel = app.storage.user.get(obdobi_klic)
    if obdobi_sel not in obdobi_klice:
        obdobi_sel = obdobi_klice[0] if obdobi_klice else None
        app.storage.user[obdobi_klic] = obdobi_sel
    radky, celkem = nacti_data_nahled(pobocka, obdobi_sel)

    async def _zpracuj(nazev: str):
        stav_pokrok = {'f': 0.0}
        dlg, kruh, popisek = _prekryv_kolecko(f'Zpracovávám {nazev}…')
        dlg.open()

        def _tik():
            proc = stav_pokrok['f'] * 100
            kruh.set_value(proc)
            popisek.set_text(f'{proc:.0f} %')

        casovac = ui.timer(0.2, _tik)
        try:
            pocet, cil, obdobi_sou = await asyncio.to_thread(
                zpracuj_soubor, pobocka, nazev, lambda f: stav_pokrok.update(f=f))
        except Exception as exc:
            ui.notify(f'Chyba zpracování: {exc}', type='negative', timeout=10000)
            return
        finally:
            casovac.cancel()
            dlg.close()
        ui.notify(f'{nazev}: {pocet} řádků → {os.path.basename(cil)}', type='positive')
        intranet_logger.log_activity(
            user_name, 'Bonusy AO',
            f'{pobocka}: zpracován soubor {nazev} ({pocet} řádků)')
        app.storage.user[obdobi_klic] = obdobi_sou
        _data_sekce.refresh()

    async def _zmen_obdobi(e):
        app.storage.user[obdobi_klic] = e.value
        dlg, _kruh, _popisek = _prekryv_kolecko(f'Načítám období {e.value}…',
                                                procenta=False)
        dlg.open()
        await asyncio.sleep(0.15)   # ať se překryv vykreslí před načtením z DB
        dlg.close()
        _data_sekce.refresh()

    async def _smaz_obdobi():
        if not obdobi_sel:
            return
        dlg_smaz.close()
        dlg, _kruh, _popisek = _prekryv_kolecko(f'Mažu období {obdobi_sel}…',
                                                procenta=False)
        dlg.open()
        try:
            pocet = await asyncio.to_thread(smaz_obdobi, pobocka, obdobi_sel)
        except Exception as exc:
            ui.notify(f'Chyba mazání: {exc}', type='negative', timeout=10000)
            return
        finally:
            dlg.close()
        ui.notify(f'Smazáno období {obdobi_sel}: {pocet} řádků.', type='positive')
        intranet_logger.log_activity(
            user_name, 'Bonusy AO',
            f'{pobocka}: smazáno období {obdobi_sel} ({pocet} řádků)')
        app.storage.user.pop(obdobi_klic, None)
        _data_sekce.refresh()

    with ui.dialog() as dlg_smaz, ui.card().classes('w-[420px]'):
        ui.label('Smazat období').classes('text-lg font-bold')
        ui.label(f'Opravdu smazat všechna data pobočky {pobocka} za období '
                 f'{obdobi_sel}? ({celkem} řádků) Akce je nevratná – soubory ve '
                 f'složce Zpracovano zůstanou.').classes('text-sm text-gray-600')
        with ui.row().classes('w-full justify-end gap-2'):
            ui.button('Zrušit', on_click=dlg_smaz.close).props('flat no-caps')
            ui.button('Smazat', icon='delete', on_click=_smaz_obdobi) \
                .props('color=negative no-caps')

    async def _on_kontakty(e):
        zdroj = None
        for att in ('content', 'file', 'stream', 'data', 'file_obj'):
            val = getattr(e, att, None)
            if val is not None and hasattr(val, 'read'):
                zdroj = val
                break
        if zdroj is None:
            ui.notify('Nepodařilo se načíst obsah souboru.', type='negative')
            return
        raw = zdroj.read()
        if inspect.isawaitable(raw):
            raw = await raw
        try:
            dvojice = await asyncio.to_thread(parsuj_kontakty, raw, e.name)
            pocet = await asyncio.to_thread(uloz_kontakty, dvojice)
        except Exception as exc:
            ui.notify(f'Chyba: {exc}', type='negative', timeout=10000)
            return
        ui.notify(f'Načteno {pocet} kontaktů.', type='positive')
        dlg_kontakty.close()
        _data_sekce.refresh()

    with ui.dialog() as dlg_kontakty, ui.card().classes('w-[560px]'):
        ui.label('Kontaktní údaje VO').classes('text-lg font-bold')
        ui.label('Nahraj nejnovější soubor „Kontaktní údaje – Změna OZ k DD.MM.YYYY“ '
                 '(.xlsb / .xlsx). Podle IČO se z něj doplňuje Jméno a K. jméno.') \
            .classes('text-xs text-gray-500')
        ui.label(f'V databázi: {pocet_kontaktu()} záznamů').classes('text-sm')
        ui.upload(on_upload=_on_kontakty, auto_upload=True, max_file_size=50_000_000,
                  label='Vybrat soubor').props('accept=.xlsb,.xlsx').classes('w-full')
        with ui.row().classes('w-full justify-end'):
            ui.button('Zavřít', on_click=dlg_kontakty.close).props('flat no-caps')

    with ui.column().classes('w-full gap-2 mt-6'):
        with ui.row().classes('w-full items-center justify-between'):
            ui.label('Data').classes('text-lg font-bold text-gray-800')
            with ui.row().classes('gap-2'):
                ui.button('Kontaktní údaje VO', icon='contacts',
                          on_click=dlg_kontakty.open) \
                    .props('color=primary outline no-caps dense')
                ui.button('Obnovit', icon='refresh', on_click=_data_sekce.refresh) \
                    .props('color=teal outline no-caps dense')
        ui.label(f'Vstup: {os.path.abspath(edit)}   •   Výstup: {os.path.abspath(hotovo)}') \
            .classes('text-xs text-gray-500')

        with ui.row().classes('w-full gap-4 items-stretch'):
            with ui.card().classes('flex-1 p-3 gap-2'):
                ui.label('Ke zpracování – složka Edit') \
                    .classes('text-sm font-semibold text-gray-700')
                soubory = _soubory(edit)
                if not soubory:
                    ui.label('Složka je prázdná.').classes('text-xs text-gray-400')
                else:
                    with ui.row().classes('items-center gap-2 w-full no-wrap'):
                        vyber = ui.select(
                            {n: f'{n} ({_velikost(v)})' for n, v in soubory},
                            value=soubory[0][0], label='Soubor') \
                            .props('dense outlined options-dense').classes('flex-grow')
                        async def _klik(_e=None, v=vyber):
                            if not v.value:
                                ui.notify('Vyberte soubor.', type='warning')
                                return
                            await _zpracuj(v.value)

                        ui.button('Zpracovat', icon='play_arrow', on_click=_klik) \
                            .props('color=primary outline no-caps dense')
            with ui.card().classes('flex-1 p-3 gap-2'):
                ui.label('Hotové – složka Zpracovano') \
                    .classes('text-sm font-semibold text-gray-700')
                hotove = _soubory(hotovo)
                if not hotove:
                    ui.label('Zatím nic zpracovaného.').classes('text-xs text-gray-400')
                for nazev, vel in hotove:
                    with ui.row().classes('items-center gap-2 w-full'):
                        ui.icon('task_alt', color='green-6')
                        ui.label(f'{nazev} ({_velikost(vel)})').classes('text-sm')
                        ui.space()
                        ui.button('Stáhnout', icon='download',
                                  on_click=lambda _e, n=nazev:
                                      ui.download.file(os.path.join(hotovo, n), n)) \
                            .props('color=green outline no-caps dense')

        with ui.row().classes('w-full items-center gap-3 mt-2'):
            if obdobi_klice:
                ui.select(obdobi_klice,
                          value=obdobi_sel, label='Období', on_change=_zmen_obdobi) \
                    .props('dense outlined options-dense').style('min-width: 220px')
                ui.button('Smazat období', icon='delete', on_click=dlg_smaz.open) \
                    .props('color=negative outline no-caps dense')
            else:
                ui.label('Žádná data v databázi.').classes('text-sm text-gray-400')
            ui.space()
            ui.label(f'V databázi: {celkem} řádků • náhled prvních {len(radky)}') \
                .classes('text-xs text-gray-500')
        ui.aggrid({
            'columnDefs': _data_col_defs(),
            'rowData': radky,
            'defaultColDef': {'resizable': True, 'sortable': True, 'filter': True},
            'rowHeight': 32,
            'suppressMovableColumns': True,
            ':onFirstDataRendered': _AUTOSIZE_FIT,
            ':onGridSizeChanged': _AUTOSIZE_FIT,
        }).classes('w-full').style(_GRID_STYLE)


def vykresli_bonusy_ao(user_id: int, user_name: str, vsechna_prava: list):
    if 'vse' not in vsechna_prava:
        with ui.column().classes('items-center py-24 gap-4'):
            ui.icon('lock', size='4rem', color='grey-4')
            ui.label('Modul Bonusy AO je zatím jen pro administrátora.') \
                .classes('text-gray-400 text-lg')
        return

    inicializace_bonusy_ao_db()
    _state_key = f'bonusy_ao_pobocka_{user_id}'

    @ui.refreshable
    def _panel():
        sel = app.storage.user.get(_state_key)
        if sel not in POBOCKY:
            ui.label('Bonusy AO').classes('text-3xl font-bold text-gray-800 mb-6')
            with ui.grid(columns=4).classes('w-full gap-4'):
                for p in POBOCKY:
                    with ui.card().classes(
                        'cursor-pointer hover:shadow-xl hover:-translate-y-0.5 transition-all '
                        'duration-200 p-6 items-center gap-2 rounded-2xl border border-gray-100'
                    ) as card:
                        ui.label('💰').classes('text-5xl mb-1')
                        ui.label(p).classes('text-lg font-bold text-center text-gray-800')
                        card.on('click', lambda pb=p: _otevri(pb))
        else:
            with ui.column().classes('w-full gap-0'):
                with ui.row().classes('items-center gap-3 mb-4'):
                    ui.button(icon='arrow_back', on_click=_zpet).props('flat round') \
                        .tooltip('Zpět na výběr pobočky')
                    ui.label(f'Bonusy AO – {sel}').classes('text-2xl font-bold text-gray-800')
                _sekce_defs = {
                    'podklad': ('📋 Podkladová tabulka',
                                lambda: _podkladova_tabulka(sel, user_name)),
                    'data':    ('🗂️ Data',
                                lambda: _data_sekce(sel, user_name)),
                    'bonusy':  ('💰 Bonusy',
                                lambda: _bonusy_sekce(sel, user_name)),
                }
                with ui.tabs().props(
                    'align=left active-color=primary indicator-color=primary'
                ).classes('w-full border-b border-gray-200') as _tabs_sekce:
                    for _k, (_lbl, _) in _sekce_defs.items():
                        ui.tab(_k, label=_lbl)
                with ui.tab_panels(_tabs_sekce, value='podklad').classes('w-full pt-4'):
                    for _k, (_, _fn) in _sekce_defs.items():
                        with ui.tab_panel(_k):
                            _fn()

    def _otevri(pb: str):
        app.storage.user[_state_key] = pb
        _panel.refresh()

    def _zpet():
        app.storage.user.pop(_state_key, None)
        _panel.refresh()

    _panel()
