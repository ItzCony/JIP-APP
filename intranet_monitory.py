"""Modul Monitor — sledování cen konkurence ve čtyřech sestavách:

  • KUPI   (Kupi …-vše.xlsm,               list „porovnání")
  • MAKRO  (Makro-…-vše.xlsm,              listy „makro-akce" + „makro-ceny")
  • VO     (Monitor VO letáky …-vše.xlsm,  list „porovnání")
  • TAMDA  (Porovnání tamda …-vše.xlsb,    list „porovnání")

Princip je stejný jako u Sankcí → „Nedodávky dod. k vyjádření": nahraje se
týdenní sestava, nákupčí k řádkům píší vyjádření, kdokoliv s právem čtení
exportuje do XLSX. Data přibývají po týdnech, období nese každý řádek a
re-import stejného období dávku nahradí — vyjádření se zachovají (párují se
přes row_hash).

Rozdíl proti sankcím: hlavička každého monitoru je jiná (30–45 sloupců) a v
čase se mění. Ruční mapování sloupec→DB by tedy znamenalo čtyři schémata, která
se rozbijí při první úpravě exportu. Proto se celý řádek ukládá do JSON sloupce
`data` a sloupce gridu se berou z hlavičky POSLEDNÍHO importu (tabulka
`monitor_meta`). Přidaný/odebraný sloupec v souboru tak projde bez migrace.

Role (intranet_prava.py), pro každý monitor zvlášť:
  • monitor_<typ>_ctenar  – vidí sestavu, filtruje, exportuje.
  • monitor_<typ>_zadatel – navíc píše vyjádření.
  • monitor_<typ>_admin   – navíc import dat a mazání sestavy.
  • 'vse'                 – vše ve všech monitorech.
"""

from nicegui import ui, app
import intranet_data
import intranet_logger
import intranet_sankce
from intranet_ui_utils import prekryv_kolecko, refreshable_na_klienta
import datetime
import hashlib
import inspect
import io
import re
import json
import unicodedata
import asyncio
from collections import OrderedDict

# Sdílené helpery ze Sankcí — tenhle modul je jejich mladší sourozenec a nemá
# smysl je psát podruhé (čtení sešitu, export XLSX, audit, historie řádku).
_norm = intranet_sankce._norm
_f = intranet_sankce._f
_s = intranet_sankce._s
_cti_list_data = intranet_sankce._cti_list_data
_parse_obdobi_z_nazvu = intranet_sankce._parse_obdobi_z_nazvu
_obdobi_label = intranet_sankce._obdobi_label
_export_xlsx = intranet_sankce._export_xlsx
_zobraz_historii = intranet_sankce._zobraz_historii
zapis_audit = intranet_sankce.zapis_audit

_META_TABULKA = 'monitor_meta'
_SABLONA_TABULKA = 'monitor_sablona'
_SORT_TABULKA = intranet_data.MONITOR_SORT_TABULKA

# Sloupec s nákupním sortimentem. Listy ho pojmenovaly různě („nák-sort"
# v Kupi/VO/Tamda/makro-akce, „nak" v makro-ceny) — bereme první, který
# v hlavičce listu je. Kódy jsou napříč monitory stejné (CK, NP, DR, …),
# proto je i právo jedno společné, ne čtvero per monitor.
_SORT_KLICE = ('nak_sort', 'nak')
# Prázdná hodnota a #N/A = nezařazené zboží. Vidí ho jen správce monitoru.
_SORT_PRAZDNE = {'', '#n/a', 'n/a', '#nd', 'none', '-'}

# ── Konfigurace monitorů ─────────────────────────────────────
# `listy`  – názvy listů v sešitu (porovnává se bez diakritiky a velikosti
#            písmen); importují se všechny, které v souboru jsou.
# `klic`   – sloupce, ze kterých se skládá row_hash (stabilní identita řádku
#            napříč re-importem téhož období). Berou se jen ty, které v hlavičce
#            opravdu jsou; když není žádný, hashuje se celý řádek.
MONITORY = OrderedDict([
    ('kupi', {
        'nazev': 'KUPI', 'emoji': '🛒', 'tabulka': 'monitor_kupi',
        'listy': ('porovnání',), 'klic': ('kod', 'k2', 'id', 'dodavatel'),
        'border': 'border-indigo-200', 'btn': 'bg-indigo-600 hover:bg-indigo-700',
    }),
    ('makro', {
        'nazev': 'MAKRO', 'emoji': '🏬', 'tabulka': 'monitor_makro',
        'listy': ('makro-akce', 'makro-ceny'), 'klic': ('kod', 'k2', 'ean', 'id'),
        'border': 'border-amber-200', 'btn': 'bg-amber-600 hover:bg-amber-700',
    }),
    ('vo', {
        'nazev': 'VO letáky', 'emoji': '📰', 'tabulka': 'monitor_vo',
        'listy': ('porovnání',), 'klic': ('kod', 'k2', 'id', 'dodavatel'),
        'border': 'border-emerald-200', 'btn': 'bg-emerald-600 hover:bg-emerald-700',
    }),
    ('tamda', {
        'nazev': 'TAMDA', 'emoji': '🛍️', 'tabulka': 'monitor_tamda',
        'listy': ('porovnání',), 'klic': ('id', 'kod_jip', 'k2'),
        'border': 'border-sky-200', 'btn': 'bg-sky-600 hover:bg-sky-700',
    }),
])

_VSE_OBD = '(všechna období)'


# ============================================================
# ==                  POMOCNÉ FUNKCE                        ==
# ============================================================

def _klic_sloupce(hlavicka) -> list:
    """Hlavička souboru → klíče pro JSON/grid. Názvy se normalizují (bez
    diakritiky, malá písmena, `_` místo mezer); duplicitní hlavičky — a ty mají
    všechny čtyři sestavy (Makro 2× „Kód", Tamda 2× „ID") — dostanou sufix _2."""
    klice, videno = [], {}
    for i, h in enumerate(hlavicka):
        txt = unicodedata.normalize('NFKD', _s(h).strip())
        txt = ''.join(c for c in txt if not unicodedata.combining(c)).lower()
        zaklad = ''.join(c if c.isalnum() else '_' for c in txt).strip('_')
        while '__' in zaklad:
            zaklad = zaklad.replace('__', '_')
        if not zaklad:
            zaklad = f'sl{i + 1}'
        videno[zaklad] = videno.get(zaklad, 0) + 1
        klice.append(zaklad if videno[zaklad] == 1 else f'{zaklad}_{videno[zaklad]}')
    return klice


def _row_hash(tabulka: str, list_nazev: str, obdobi: str, casti: list, poradi: int = 1) -> str:
    """Stabilní otisk řádku v rámci období — přes něj se při re-importu téhož
    období párují ručně zapsaná vyjádření.

    `poradi` je pořadí stejného klíče v souboru: totéž zboží je v monitorech
    běžně na víc řádcích (jiná prodejna, jiná akce), takže samotný klíč
    (Kód/K2/…) unikátní není a bez rozlišení by se řádky slévaly."""
    klic = '|'.join([tabulka, list_nazev, obdobi] + [_s(c) for c in casti] + [f'#{poradi}'])
    return hashlib.md5(klic.encode('utf-8')).hexdigest()


def _obdobi_z_nazvu(nazev: str) -> tuple:
    """Období z názvu souboru monitoru. Sankce mají v názvu rozsah „01.05 -15.05.2026",
    monitory jedno datum („Kupi 2026-09-16-vše") nebo měsíc („…VO letáky 2026-09-vše").
    Vrací (od_iso, do_iso) nebo (None, None)."""
    od, do = _parse_obdobi_z_nazvu(nazev)        # nejdřív formát ze sankcí
    if od and do:
        return od, do
    if not nazev:
        return None, None
    m = re.search(r'(20\d{2})[-_.](\d{1,2})[-_.](\d{1,2})', nazev)
    if m:
        try:
            d = datetime.date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
            return d.isoformat(), d.isoformat()
        except ValueError:
            pass
    m = re.search(r'(\d{1,2})\.(\d{1,2})\.(20\d{2})', nazev)
    if m:
        try:
            d = datetime.date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
            return d.isoformat(), d.isoformat()
        except ValueError:
            pass
    m = re.search(r'(20\d{2})[-_.](\d{1,2})(?!\d)', nazev)     # celý měsíc
    if m:
        try:
            od = datetime.date(int(m.group(1)), int(m.group(2)), 1)
            do = (datetime.date(od.year + (od.month == 12), od.month % 12 + 1, 1)
                  - datetime.timedelta(days=1))
            return od.isoformat(), do.isoformat()
        except ValueError:
            pass
    return None, None


# ============================================================
# ==      ŠABLONA SEŠITU (vzorce, formáty, CF, šířky)       ==
# ============================================================
# Export má být 1:1 vůči zdroji. Nejlevnější cesta, jak toho dosáhnout, je
# nechat originál dělat práci za nás: při importu se sešit ořízne na hlavičku
# + první datový řádek a uloží se do `monitor_meta.sablona`. Ten zbytek nese
# všechno podstatné — styly, vzorce, podmíněné formátování, šířky sloupců,
# příčky. Export pak jen dosype řádky a vzorce přeloží na jejich čísla řádků.
# Zároveň si vytáhnu formáty čísel a CF pravidla pro zobrazení v gridu.
_SABLONA_MAX = 8 * 1024 * 1024        # pojistka proti obřím sešitům v DB
_CF_ZNAK = {'lessThan': '<', 'lessThanOrEqual': '<=', 'greaterThan': '>',
            'greaterThanOrEqual': '>=', 'equal': '==', 'notEqual': '!='}


def _barva(o):
    """openpyxl Color → '#RRGGBB', nebo None. Color.rgb vrací u theme/auto
    barev místo None hlášku descriptoru ("Values must be of type <class
    'str'>"). Ta je pravdivostně true, takže se z ní dřív skládalo CSS
    `#'str'>` a prohlížeč pravidlo zahodil. Beru jen skutečný hex — theme
    barvy grid nekreslí, v exportu zůstanou přes šablonu."""
    rgb = getattr(o, 'rgb', None) if o is not None else None
    if not isinstance(rgb, str) or len(rgb) not in (6, 8):
        return None
    hexa = rgb[-6:]
    try:
        int(hexa, 16)
    except ValueError:
        return None
    return '#' + hexa


def _bunka_styl(bunka) -> dict:
    """Statický styl buňky ze sešitu → CSS dict pro grid. Bílou výplň a černé
    písmo vynechávám — to grid kreslí tak jako tak a jen by to nafouklo meta."""
    s = {}
    f = bunka.fill
    if f is not None and f.patternType == 'solid':
        bg = _barva(f.fgColor)
        if bg and bg != '#FFFFFF':
            s['backgroundColor'] = bg
    font = bunka.font
    if font is not None:
        if font.b:
            s['fontWeight'] = 'bold'
        if font.i:
            s['fontStyle'] = 'italic'
        fg = _barva(font.color)
        if fg and fg != '#000000':
            s['color'] = fg
    al = bunka.alignment.horizontal if bunka.alignment is not None else None
    if al in ('left', 'center', 'right'):
        s['textAlign'] = al
    return s


def _cf_pravidlo(r):
    """CF pravidlo z openpyxl → jednoduchý dict pro grid, nebo None (dataBary,
    barevné škály a pravidla bez viditelného formátu grid nekreslí — v exportu
    zůstanou, ty jedou přes šablonu)."""

    dxf = getattr(r, 'dxf', None)
    if dxf is None:
        return None
    bg = _barva(dxf.fill.bgColor) if dxf.fill is not None else None
    fg = _barva(dxf.font.color) if dxf.font is not None else None
    tucne = bool(dxf.font.b) if dxf.font is not None else False
    if not (bg or fg or tucne):
        return None
    p = {'bg': bg, 'fg': fg, 'b': tucne}
    if r.type == 'containsText' and r.text:
        return dict(p, op='contains', v=r.text)
    if r.type == 'cellIs' and r.formula:
        hod = [_f(x) for x in list(r.formula)[:2]]
        if hod[0] is None:
            return None
        return dict(p, op=r.operator, v=hod[0],
                    v2=hod[1] if len(hod) > 1 else None)
    return None


def _sablona_a_styly(raw: bytes):
    """(bajty šablony | None, {norm(list): {'fmt': {písmeno: formát},
    'cf': {písmeno: [pravidla]}, 'hdr': {písmeno: css}, 'bunka': {písmeno:
    css}}}). Jen pro .xlsx/.xlsm — .xlsb styly nenese."""
    try:
        import openpyxl
        from openpyxl.utils import get_column_letter, range_boundaries
        wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=False)
    except Exception as e:
        print(f'[monitory] Šablonu nelze načíst ({e}) — export bude bez formátů.')
        return None, {}

    styly = {}
    for ws in wb.worksheets:
        fmt, cf, hdr, bunky = {}, {}, {}, {}
        for c in range(1, (ws.max_column or 0) + 1):
            pismeno = get_column_letter(c)
            f = ws.cell(2, c).number_format
            if f and f != 'General':
                fmt[pismeno] = f
            # Hlavička z řádku 1, statický styl dat ze vzorového řádku 2.
            h = _bunka_styl(ws.cell(1, c))
            if h:
                hdr[pismeno] = h
            b = _bunka_styl(ws.cell(2, c))
            if b:
                bunky[pismeno] = b
        for rng in ws.conditional_formatting:
            for r in rng.rules:
                pravidlo = _cf_pravidlo(r)
                if not pravidlo:
                    continue
                for sq in str(rng.sqref).split():
                    c1, _, c2, _ = range_boundaries(sq)
                    for c in range(c1, c2 + 1):
                        cf.setdefault(get_column_letter(c), []).append(pravidlo)
        styly[_norm(ws.title)] = {'fmt': fmt, 'cf': cf, 'hdr': hdr,
                                  'bunka': bunky}

        # Ořez na hlavičku + vzorový řádek. `_cells` sahám záměrně přímo —
        # delete_rows() přeskládává buňky po jedné a u 20 tisíc řádků to trvá
        # minuty, tohle je konstantní.
        ws._cells = {k: v for k, v in ws._cells.items() if k[0] <= 2}
        ws._current_row = 2
        for r in [x for x in ws.row_dimensions if x > 2]:
            del ws.row_dimensions[r]

    try:
        bio = io.BytesIO()
        wb.save(bio)
        data = bio.getvalue()
    except Exception as e:
        print(f'[monitory] Šablonu nelze uložit: {e}')
        return None, styly
    if len(data) > _SABLONA_MAX:
        print(f'[monitory] Šablona má {len(data)} B — neukládám.')
        return None, styly
    return data, styly


def inicializace_monitor_db():
    """Čtyři tabulky monitorů + tabulka hlaviček. Volá se při startu."""
    conn = intranet_data.get_db_connection()
    if not conn:
        return
    try:
        cur = conn.cursor()
        for spec in MONITORY.values():
            cur.execute(f"""
                CREATE TABLE IF NOT EXISTS {spec['tabulka']} (
                    id INT AUTO_INCREMENT PRIMARY KEY,
                    obdobi VARCHAR(60),
                    obdobi_od DATE, obdobi_do DATE,
                    list_nazev VARCHAR(60) NOT NULL DEFAULT '',
                    row_hash VARCHAR(40),
                    data JSON NOT NULL,
                    vyjadreni TEXT,
                    vyjadreni_by VARCHAR(255),
                    vyjadreni_at DATETIME DEFAULT NULL,
                    import_at DATETIME DEFAULT CURRENT_TIMESTAMP,
                    imported_by VARCHAR(255),
                    INDEX idx_obdobi (obdobi), INDEX idx_hash (row_hash),
                    INDEX idx_list (list_nazev)
                ) CHARACTER SET utf8mb4 COLLATE utf8mb4_unicode_ci
            """)
        # Hlavičky: co který list naposledy přinesl (pořadí + popisky + čísla).
        cur.execute(f"""
            CREATE TABLE IF NOT EXISTS {_META_TABULKA} (
                tabulka VARCHAR(40) NOT NULL,
                list_nazev VARCHAR(60) NOT NULL,
                sloupce JSON NOT NULL,
                updated_at DATETIME DEFAULT CURRENT_TIMESTAMP
                                    ON UPDATE CURRENT_TIMESTAMP,
                PRIMARY KEY (tabulka, list_nazev)
            ) CHARACTER SET utf8mb4 COLLATE utf8mb4_unicode_ci
        """)
        cur.execute("SELECT COUNT(*) FROM information_schema.COLUMNS WHERE "
                    "TABLE_SCHEMA=DATABASE() AND TABLE_NAME=%s AND "
                    "COLUMN_NAME='styly'", (_META_TABULKA,))
        if cur.fetchone()[0] == 0:
            cur.execute(f'ALTER TABLE {_META_TABULKA} ADD COLUMN styly JSON')
        # Prázdný sešit ze zdroje (hlavička + vzorový řádek) pro export 1:1.
        cur.execute(f"""
            CREATE TABLE IF NOT EXISTS {_SABLONA_TABULKA} (
                tabulka VARCHAR(40) NOT NULL PRIMARY KEY,
                pripona VARCHAR(10) NOT NULL,
                data LONGBLOB NOT NULL,
                updated_at DATETIME DEFAULT CURRENT_TIMESTAMP
                                    ON UPDATE CURRENT_TIMESTAMP
            ) CHARACTER SET utf8mb4 COLLATE utf8mb4_unicode_ci
        """)
        # Číselník nákupních sortimentů — zdroj práv „monitor_sort_*".
        # Kódy se jen přidávají, nikdy nemažou (viz ziskej_sortimenty_monitoru).
        cur.execute(f"""
            CREATE TABLE IF NOT EXISTS {_SORT_TABULKA} (
                kod VARCHAR(20) NOT NULL PRIMARY KEY,
                popis VARCHAR(80) NOT NULL DEFAULT '',
                prvni_videno DATETIME DEFAULT CURRENT_TIMESTAMP,
                posledni_videno DATETIME DEFAULT CURRENT_TIMESTAMP
                                    ON UPDATE CURRENT_TIMESTAMP
            ) CHARACTER SET utf8mb4 COLLATE utf8mb4_unicode_ci
        """)
        conn.commit()
        cur.execute(f'SELECT COUNT(*) FROM {_SORT_TABULKA}')
        prazdny = cur.fetchone()[0] == 0
        cur.close()
        if prazdny:
            _backfill_sortimenty(conn)
    except Exception as e:
        print(f'Chyba při inicializaci DB Monitorů: {e}')
    finally:
        conn.close()


def _backfill_sortimenty(conn):
    """Jednorázově: kódy z už naimportovaných dat, ať se nemusí reimportovat."""
    try:
        cur = conn.cursor()
        cur.execute(f'SELECT tabulka, list_nazev, sloupce FROM {_META_TABULKA}')
        listy = cur.fetchall()
        kody = set()
        for tabulka, list_nazev, sloupce in listy:
            try:
                sl = json.loads(sloupce) if isinstance(sloupce, str) else (sloupce or [])
            except Exception:
                continue
            klic = _sort_klic(sl)
            if not klic:
                continue
            cur.execute(
                f"SELECT DISTINCT JSON_UNQUOTE(JSON_EXTRACT(data, %s)) "
                f"FROM {tabulka} WHERE list_nazev=%s", (f'$.{klic}', list_nazev))
            kody |= {_sort_kod(v) for (v,) in cur.fetchall()}
        cur.close()
        kody = {k for k in kody if k}
        if not kody:
            return
        cur = conn.cursor()
        cur.executemany(f'INSERT IGNORE INTO {_SORT_TABULKA} (kod) VALUES (%s)',
                        [(k,) for k in sorted(kody)])
        conn.commit()
        cur.close()
        intranet_data.zneplatni_cache_sortimentu()
        print(f'[monitory] Číselník sortimentů doplněn z dat: {len(kody)} kódů.')
    except Exception as e:
        print(f'[monitory] Backfill sortimentů selhal: {e}')


def _uloz_meta(tabulka: str, list_nazev: str, sloupce: list, styly: dict = None):
    conn = intranet_data.get_db_connection()
    if not conn:
        return
    try:
        cur = conn.cursor()
        cur.execute(
            f'INSERT INTO {_META_TABULKA} (tabulka,list_nazev,sloupce,styly) '
            f'VALUES (%s,%s,%s,%s) '
            f'ON DUPLICATE KEY UPDATE sloupce=VALUES(sloupce), styly=VALUES(styly)',
            (tabulka, list_nazev, json.dumps(sloupce, ensure_ascii=False),
             json.dumps(styly, ensure_ascii=False) if styly else None))
        conn.commit()
        cur.close()
    finally:
        conn.close()


def _uloz_sablonu(tabulka: str, pripona: str, data: bytes):
    conn = intranet_data.get_db_connection()
    if not conn:
        return
    try:
        cur = conn.cursor()
        cur.execute(
            f'INSERT INTO {_SABLONA_TABULKA} (tabulka,pripona,data) VALUES (%s,%s,%s) '
            f'ON DUPLICATE KEY UPDATE pripona=VALUES(pripona), data=VALUES(data)',
            (tabulka, pripona, data))
        conn.commit()
        cur.close()
    finally:
        conn.close()


def _nacti_sablonu(tabulka: str):
    """(bajty, přípona) posledního importovaného sešitu, nebo (None, None)."""
    conn = intranet_data.get_db_connection()
    if not conn:
        return None, None
    try:
        cur = conn.cursor()
        cur.execute(f'SELECT data, pripona FROM {_SABLONA_TABULKA} WHERE tabulka=%s',
                    (tabulka,))
        r = cur.fetchone()
        cur.close()
        return (r[0], r[1]) if r else (None, None)
    finally:
        conn.close()


def _nacti_styly(tabulka: str) -> dict:
    """{list_nazev: {'fmt': {…}, 'cf': {…}}} z posledního importu."""
    conn = intranet_data.get_db_connection()
    if not conn:
        return {}
    try:
        cur = conn.cursor()
        cur.execute(f'SELECT list_nazev, styly FROM {_META_TABULKA} WHERE tabulka=%s',
                    (tabulka,))
        out = {}
        for ln, st in cur.fetchall():
            if not st:
                continue
            try:
                out[ln] = json.loads(st) if isinstance(st, str) else st
            except Exception:
                pass
        cur.close()
        return out
    finally:
        conn.close()


def _nacti_meta(tabulka: str) -> dict:
    """{list_nazev: [{'k','h','num'}, …]} podle posledního importu."""
    conn = intranet_data.get_db_connection()
    if not conn:
        return {}
    try:
        cur = conn.cursor()
        cur.execute(f'SELECT list_nazev, sloupce FROM {_META_TABULKA} WHERE tabulka=%s',
                    (tabulka,))
        out = {}
        for ln, sl in cur.fetchall():
            try:
                out[ln] = json.loads(sl) if isinstance(sl, str) else sl
            except Exception:
                pass
        cur.close()
        return out
    finally:
        conn.close()


def _nacti(tabulka: str, list_nazev: str = None) -> list:
    """Řádky sestavy: JSON `data` rozbalené do plochého dictu pro grid."""
    conn = intranet_data.get_db_connection()
    if not conn:
        return []
    try:
        cur = conn.cursor(dictionary=True)
        sql = (f'SELECT id,obdobi,list_nazev,row_hash,data,vyjadreni,vyjadreni_by,'
               f'vyjadreni_at FROM {tabulka}')
        params = ()
        if list_nazev:
            sql += ' WHERE list_nazev=%s'
            params = (list_nazev,)
        sql += ' ORDER BY obdobi_od DESC, id ASC'
        cur.execute(sql, params)
        rows = []
        for r in cur.fetchall():
            try:
                data = json.loads(r['data']) if isinstance(r['data'], str) else (r['data'] or {})
            except Exception:
                data = {}
            radek = dict(data)
            radek.update({
                'id': r['id'], 'obdobi': r['obdobi'], 'list_nazev': r['list_nazev'],
                'row_hash': r['row_hash'], 'vyjadreni': r['vyjadreni'] or '',
                'vyjadreni_by': r['vyjadreni_by'] or '',
                'vyjadreni_at': r['vyjadreni_at'].strftime('%d.%m.%Y %H:%M')
                                if r['vyjadreni_at'] else '',
            })
            rows.append(radek)
        cur.close()
        return rows
    finally:
        conn.close()


def _seznam_obdobi(tabulka: str, list_nazev: str = None) -> list:
    conn = intranet_data.get_db_connection()
    if not conn:
        return []
    try:
        cur = conn.cursor()
        sql = f'SELECT obdobi FROM {tabulka} WHERE obdobi IS NOT NULL'
        params = ()
        if list_nazev:
            sql += ' AND list_nazev=%s'
            params = (list_nazev,)
        cur.execute(sql + ' GROUP BY obdobi, obdobi_od ORDER BY obdobi_od DESC', params)
        return [r[0] for r in cur.fetchall() if r[0]]
    finally:
        conn.close()


# ============================================================
# ==                       IMPORT                           ==
# ============================================================

def _importuj_sync(raw: bytes, typ: str, obdobi: str, od_iso: str, do_iso: str,
                   user_name: str, pokrok=None, nazev_souboru: str = '') -> tuple:
    """Naimportuje VŠECHNY listy monitoru ze sešitu. Dávku daného období
    nahradí, vyjádření se zachovají (párování přes row_hash).
    Vrací (počet_řádků, seznam_listů, chyba|None)."""
    spec = MONITORY[typ]
    tabulka = spec['tabulka']
    celkem, hotove = 0, []

    # Šablona pro export 1:1. .xlsb openpyxl neotevře a pyxlsb styly nečte —
    # tam zůstane export bez formátů (viz _export_monitor_xlsx).
    pripona = (nazev_souboru.rsplit('.', 1)[-1] or '').lower()
    styly_listu = {}
    if pripona in ('xlsx', 'xlsm'):
        sablona, styly_listu = _sablona_a_styly(raw)
        if sablona:
            _uloz_sablonu(tabulka, pripona, sablona)

    conn = intranet_data.get_db_connection()
    if not conn:
        return 0, [], 'Chyba připojení k databázi.'
    try:
        for idx, nazev_listu in enumerate(spec['listy']):
            def _dil(x, _i=idx):
                if pokrok:
                    pokrok((_i + x) / max(len(spec['listy']), 1))

            header, rows_iter, chyba = _cti_list_data(raw, _dil, list_nazev=nazev_listu)
            if chyba:
                # Makro má dva listy; když v souboru jeden chybí, není to chyba
                # celého importu — jen se nenaimportuje.
                if len(spec['listy']) > 1:
                    continue
                return 0, [], chyba

            klice = _klic_sloupce(header)
            popisky = [_s(h).strip() or k for h, k in zip(header, klice)]
            zaznamy, cisla = [], {k: [0, 0] for k in klice}   # [číselných, neprázdných]
            klic_idx = [i for i, k in enumerate(klice) if k in spec['klic']]
            poradi = {}      # klíč řádku → kolikátý výskyt v souboru
            k_sort = next((k for k in _SORT_KLICE if k in klice), None)
            sorty = set()    # kódy sortimentů viděné v tomhle listu

            for r in rows_iter:
                if r is None or all(c is None or c == '' for c in r):
                    continue
                data = {}
                for i, k in enumerate(klice):
                    v = r[i] if i < len(r) else None
                    if v is None or v == '':
                        continue
                    if isinstance(v, (int, float)) and not isinstance(v, bool):
                        data[k] = v
                        cisla[k][0] += 1
                        cisla[k][1] += 1
                    else:
                        txt = _s(v).strip()
                        if not txt:
                            continue
                        data[k] = txt
                        cisla[k][1] += 1
                        if _f(txt) is not None:
                            cisla[k][0] += 1
                if not data:
                    continue
                if k_sort:
                    sorty.add(_sort_kod(data.get(k_sort)))
                casti = ([r[i] if i < len(r) else '' for i in klic_idx] if klic_idx
                         else [data.get(k, '') for k in klice])
                ck = '|'.join(_s(c) for c in casti)
                poradi[ck] = poradi.get(ck, 0) + 1
                zaznamy.append((_row_hash(tabulka, nazev_listu, obdobi, casti, poradi[ck]),
                                json.dumps(data, ensure_ascii=False, default=str)))

            if not zaznamy:
                if len(spec['listy']) > 1:
                    continue
                return 0, [], f'List „{nazev_listu}" neobsahuje žádné datové řádky.'

            cur = conn.cursor(dictionary=True)
            cur.execute(f'SELECT row_hash,vyjadreni,vyjadreni_by,vyjadreni_at FROM {tabulka} '
                        f'WHERE obdobi=%s AND list_nazev=%s', (obdobi, nazev_listu))
            zachov = {x['row_hash']: x for x in cur.fetchall()}
            cur.close()

            cur2 = conn.cursor()
            cur2.execute(f'DELETE FROM {tabulka} WHERE obdobi=%s AND list_nazev=%s',
                         (obdobi, nazev_listu))
            davka = []
            for rh, js in zaznamy:
                z = zachov.get(rh, {})
                davka.append((obdobi, od_iso or None, do_iso or None, nazev_listu, rh, js,
                              z.get('vyjadreni'), z.get('vyjadreni_by'),
                              z.get('vyjadreni_at'), user_name))
            cur2.executemany(
                f'INSERT INTO {tabulka} (obdobi,obdobi_od,obdobi_do,list_nazev,row_hash,'
                f'data,vyjadreni,vyjadreni_by,vyjadreni_at,imported_by) '
                f'VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)', davka)
            conn.commit()
            cur2.close()

            # `styly_listu` je klíčované normalizovaným názvem listu v sešitu;
            # u jednolistých monitorů („DATA") se název hledá volně, tak vezmu
            # ten jediný, co tam je.
            st = styly_listu.get(_norm(nazev_listu))
            if st is None and len(styly_listu) == 1:
                st = next(iter(styly_listu.values()))
            _uloz_meta(tabulka, nazev_listu, [
                {'k': k, 'h': h, 'num': cisla[k][1] > 0 and cisla[k][0] >= cisla[k][1] * 0.8}
                for k, h in zip(klice, popisky)
            ], st)
            _uloz_sortimenty(sorty)
            celkem += len(davka)
            hotove.append(nazev_listu)

        if not hotove:
            listy = '" / „'.join(spec['listy'])
            return 0, [], f'Soubor neobsahuje list „{listy}".'
        if pokrok:
            pokrok(1.0)
        return celkem, hotove, None
    except Exception as e:
        try:
            conn.rollback()
        except Exception:
            pass
        return 0, [], f'Chyba zápisu do databáze: {e}'
    finally:
        conn.close()


def _smaz_data(tabulka: str, obdobi=None) -> tuple:
    """Nevratně smaže sestavu (celou, nebo jedno období) i s historií změn."""
    conn = intranet_data.get_db_connection()
    if not conn:
        return 0, 'Chyba připojení k databázi.'
    try:
        cur = conn.cursor()
        if obdobi is None:
            cur.execute(f'SELECT COUNT(*) FROM {tabulka}')
            pocet = cur.fetchone()[0]
            cur.execute('DELETE FROM sankce_audit WHERE tabulka=%s', (tabulka,))
            cur.execute(f'DELETE FROM {tabulka}')
        else:
            cur.execute(f'SELECT COUNT(*) FROM {tabulka} WHERE obdobi=%s', (obdobi,))
            pocet = cur.fetchone()[0]
            cur.execute(
                'DELETE FROM sankce_audit WHERE tabulka=%s AND row_hash IN '
                f'(SELECT row_hash FROM {tabulka} WHERE obdobi=%s)', (tabulka, obdobi))
            cur.execute(f'DELETE FROM {tabulka} WHERE obdobi=%s', (obdobi,))
        conn.commit()
        cur.close()
        return pocet, None
    except Exception as e:
        return 0, f'Chyba mazání: {e}'
    finally:
        conn.close()


def _uloz_vyjadreni(tabulka: str, radky: list, text: str, user_id, user_name: str) -> int:
    """Zápis vyjádření do daných řádků (DB + audit + razítko kdo/kdy)."""
    cile = [r for r in radky if (r.get('vyjadreni') or '') != (text or '')]
    if not cile:
        return 0
    kdy = datetime.datetime.now().strftime('%d.%m.%Y %H:%M')
    conn = intranet_data.get_db_connection()
    if not conn:
        ui.notify('Není spojení s databází — vyjádření se neuložilo.', type='negative')
        return 0
    try:
        cur = conn.cursor()
        cur.executemany(
            f'UPDATE {tabulka} SET vyjadreni=%s, vyjadreni_by=%s, vyjadreni_at=NOW() '
            f'WHERE id=%s', [(text, user_name, r.get('id')) for r in cile])
        conn.commit()
        cur.close()
    finally:
        conn.close()
    for r in cile:
        zapis_audit(tabulka, r.get('row_hash'), r.get('id'), 'vyjadreni',
                    r.get('vyjadreni'), text, user_id, user_name)
        r['vyjadreni'] = text
        r['vyjadreni_by'] = user_name
        r['vyjadreni_at'] = kdy
    return len(cile)


# ============================================================
# ==                        PRÁVA                           ==
# ============================================================

def prava_typu(typ: str, vsechna_prava) -> tuple:
    """(vidí, smí psát vyjádření, je správce) pro jeden monitor."""
    ma_vse = 'vse' in vsechna_prava
    admin = ma_vse or f'monitor_{typ}_admin' in vsechna_prava
    psat = admin or f'monitor_{typ}_zadatel' in vsechna_prava
    vidi = psat or f'monitor_{typ}_ctenar' in vsechna_prava
    return vidi, psat, admin


def dostupne_typy(vsechna_prava) -> list:
    return [t for t in MONITORY if prava_typu(t, vsechna_prava)[0]]


# ── Sortiment: kdo vidí které řádky ──────────────────────────
# Právo `monitor_sort_<kód>` = „vidím zboží tohoto nákupního sortimentu".
# Platí napříč všemi monitory; správce monitoru filtru nepodléhá, jinak by
# neměl jak zkontrolovat, co mu import přinesl.

_SORT_PREFIX = 'monitor_sort_'


def pravo_sortimentu(kod: str) -> str:
    return f'{_SORT_PREFIX}{_norm(kod)}'


def povolene_sorty(vsechna_prava) -> set:
    """Kódy sortimentů z práv uživatele (normalizované, malými)."""
    return {p[len(_SORT_PREFIX):] for p in vsechna_prava
            if isinstance(p, str) and p.startswith(_SORT_PREFIX)}


def _sort_klic(sloupce: list):
    """Klíč sloupce s nákupním sortimentem v tomhle listu, nebo None."""
    klice = {s.get('k') for s in sloupce}
    return next((k for k in _SORT_KLICE if k in klice), None)


def _sort_kod(hodnota) -> str:
    """Hodnota buňky → kód sortimentu; '' pro nezařazené (#N/A, prázdno)."""
    txt = _norm(_s(hodnota).strip())
    return '' if txt in _SORT_PRAZDNE else txt


def filtruj_dle_sortimentu(radky: list, sloupce: list, vsechna_prava,
                           je_admin: bool) -> list:
    """Řádky, které uživatel smí vidět. Správce dostane vše beze změny."""
    if je_admin or 'vse' in vsechna_prava:
        return radky
    klic = _sort_klic(sloupce)
    if not klic:          # list bez sloupce sortimentu — nemáme podle čeho dělit
        return radky
    povolene = povolene_sorty(vsechna_prava)
    if not povolene:
        return []
    return [r for r in radky if _sort_kod(r.get(klic)) in povolene]


def _uloz_sortimenty(kody):
    """Doplní číselník o nově viděné kódy. Existující nechá být."""
    kody = {k for k in kody if k}
    if not kody:
        return
    conn = intranet_data.get_db_connection()
    if not conn:
        return
    try:
        cur = conn.cursor()
        cur.executemany(
            f'INSERT INTO {_SORT_TABULKA} (kod) VALUES (%s) '
            f'ON DUPLICATE KEY UPDATE posledni_videno=CURRENT_TIMESTAMP',
            [(k,) for k in sorted(kody)])
        conn.commit()
        cur.close()
        intranet_data.zneplatni_cache_sortimentu()
    except Exception as e:
        print(f'[monitory] Číselník sortimentů nelze doplnit: {e}')
    finally:
        conn.close()


# ============================================================
# ==                     GRID + EXPORT                      ==
# ============================================================

def _do_sablony(raw: bytes, pripona: str, list_nazev: str, sloupce: list,
                rows: list, extra: list) -> bytes:
    """Naleje řádky do uložené šablony: zachová vzorce (přeložené na svůj řádek),
    formáty a podmíněné formátování z originálu. Vrací bajty sešitu."""
    import copy
    import openpyxl
    from openpyxl.formatting.formatting import ConditionalFormattingList
    from openpyxl.formula.translate import Translator
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.formula import ArrayFormula

    wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=False,
                                keep_vba=(pripona == 'xlsm'))
    jmeno = next((s for s in wb.sheetnames if _norm(s) == _norm(list_nazev)), None)
    ws = wb[jmeno] if jmeno else wb.worksheets[0]
    for s in wb.worksheets:
        if s is not ws:
            wb.remove(s)

    # Řádek 2 je vzor: z něj beru vzorec i styl pro každý sloupec. Maticový
    # vzorec (Kupi má takový ve sloupci „nejvýhodnější NC") nese openpyxl jako
    # ArrayFormula, ne jako řetězec — bez téhle větve by z něj v exportu zbyla
    # jen napočtená hodnota.
    def _vzor_vzorec(v):
        if isinstance(v, ArrayFormula):
            return v.text, True
        if isinstance(v, str) and v.startswith('='):
            return v, False
        return None, False

    n = len(sloupce)
    vzor = [ws.cell(2, c) for c in range(1, n + 1)]
    sablona_vzorce = [_vzor_vzorec(v.value) for v in vzor]
    styly_bunek = [v._style for v in vzor]

    for i, r in enumerate(rows):
        radek = 2 + i
        for c in range(1, n + 1):
            bunka = ws.cell(radek, c)
            f, matice = sablona_vzorce[c - 1]
            if f:
                adresa = f'{get_column_letter(c)}{radek}'
                prelozeny = Translator(f, origin=f'{get_column_letter(c)}2'
                                       ).translate_formula(adresa)
                bunka.value = ArrayFormula(adresa, prelozeny) if matice else prelozeny
            else:
                bunka.value = r.get(sloupce[c - 1]['k'])
            bunka._style = styly_bunek[c - 1]

    # Vyjádření a metadata za poslední sloupec originálu — v šabloně nejsou.
    for j, (nadpis, pole) in enumerate(extra):
        c = n + 1 + j
        ws.cell(1, c).value = nadpis
        ws.cell(1, c)._style = ws.cell(1, max(1, n))._style
        for i, r in enumerate(rows):
            ws.cell(2 + i, c).value = r.get(pole)
        ws.column_dimensions[get_column_letter(c)].width = 40 if j == 0 else 18

    zbytek = ws.max_row - (1 + len(rows))
    if zbytek > 0:   # šablona nese jen řádek 2, tohle je pojistka
        ws.delete_rows(2 + len(rows), zbytek)

    # Podmíněné formátování i autofiltr si ze šablony nesou rozsah původního
    # sešitu (u Kupi …1005). Exportovaných řádků je jiný počet — u „všech
    # období" násobně víc — takže rozsahy přetáhnu na skutečnou velikost, jinak
    # by řádky pod hranicí zůstaly bez CF. Celosloupcové rozsahy sedí samy.
    posledni = 1 + len(rows)
    prepis = []
    for rng in ws.conditional_formatting:
        casti = []
        for cr in rng.sqref.ranges:
            cr = copy.copy(cr)   # in-place by rozbilo hash v ConditionalFormattingList
            if cr.max_row < 1048576:
                cr.max_row = max(posledni, cr.min_row)
            casti.append(str(cr))
        prepis.append((' '.join(casti), list(rng.rules)))
    ws.conditional_formatting = ConditionalFormattingList()
    for sqref, pravidla in prepis:
        for pravidlo in pravidla:
            ws.conditional_formatting.add(sqref, pravidlo)

    if ws.auto_filter.ref:
        ws.auto_filter.ref = f'A1:{get_column_letter(n + len(extra))}{posledni}'

    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def _pismeno(i: int) -> str:
    """1 → 'A'. Bez importu openpyxl kvůli jedné funkci."""
    s = ''
    while i > 0:
        i, z = divmod(i - 1, 26)
        s = chr(65 + z) + s
    return s


def _fmt_js(fmt: str):
    """Excel formát čísla → JS valueFormatter. Umí to, co se ve sbírkách
    monitorů reálně vyskytuje: desetinná místa, oddělovač tisíců, procenta.
    Zbytek (datumy, měny s textem) nechá být — grid ukáže holou hodnotu, ale
    v exportu formát zůstane, ten jede přes šablonu."""
    f = fmt.split(';')[0]
    if any(z in f for z in ('y', 'd', 'h', 's')) and '0' not in f.replace('0%', ''):
        return None
    proc = '%' in f
    des = len(f.split('.')[1].split('%')[0].rstrip('_-)" ')) if '.' in f else 0
    tis = ',#' in f
    if not proc and not tis and des == 0:
        return None
    return (f"(p) => {{ const v = Number(p.value); "
            f"if (p.value === null || p.value === '' || isNaN(v)) return p.value; "
            f"return (v * {100 if proc else 1}).toLocaleString('cs-CZ', "
            f"{{minimumFractionDigits: {des}, maximumFractionDigits: {des}, "
            f"useGrouping: {'true' if tis else 'false'}}}) + '{'%' if proc else ''}'; }}")


def _cf_js(pravidla: list, zaklad: dict = None):
    """CF pravidla → JS cellStyle. První pravidlo, které sedí, vyhrává —
    stejné pořadí jako v Excelu. `zaklad` je statický styl buňky: CF v Excelu
    přepíše jen to, co dxf nese, zbytek (tučné, zarovnání) zůstává."""
    vetve = []
    for p in pravidla:
        if p.get('op') == 'contains':
            podm = f"String(p.value ?? '').includes({json.dumps(p['v'])})"
        elif p.get('op') == 'between' and p.get('v2') is not None:
            podm = f"!isNaN(v) && v >= {p['v']} && v <= {p['v2']}"
        else:
            znak = _CF_ZNAK.get(p.get('op'))
            if not znak or p.get('v') is None:
                continue
            podm = f"!isNaN(v) && v {znak} {p['v']}"
        styl = dict(zaklad or {})
        if p.get('bg'):
            styl['backgroundColor'] = p['bg']
        if p.get('fg'):
            styl['color'] = p['fg']
        if p.get('b'):
            styl['fontWeight'] = 'bold'
        vetve.append(f"if ({podm}) return {json.dumps(styl)};")
    if not vetve:
        return None
    return ("(p) => { const v = Number(p.value); " + ' '.join(vetve)
            + f" return {json.dumps(zaklad) if zaklad else 'null'}; }}")


def _col_defs(sloupce: list, editovatelne: bool, styly: dict = None) -> list:
    cols = [
        {'headerName': '', 'field': '_eye', 'width': 46, 'minWidth': 46, 'maxWidth': 46,
         'pinned': 'left', 'sortable': False, 'editable': False, 'resizable': False,
         'suppressSizeToFit': True, 'suppressAutoSize': True,
         ':cellRenderer': intranet_sankce._EYE_RENDERER,
         'cellStyle': {'textAlign': 'center', 'cursor': 'pointer', 'padding': '0'},
         'headerTooltip': 'Historie změn řádku'},
        {'headerName': 'Období', 'field': 'obdobi', 'width': 150, 'pinned': 'left',
         'sortable': True, 'cellStyle': {'fontSize': '12px', 'color': '#475569'}},
    ]
    fmt = (styly or {}).get('fmt') or {}
    cf = (styly or {}).get('cf') or {}
    hdr = (styly or {}).get('hdr') or {}
    bunky = (styly or {}).get('bunka') or {}
    for idx, s in enumerate(sloupce):
        c = {'headerName': s['h'], 'field': s['k'], 'width': 150, 'sortable': True}
        if s.get('num'):
            c['type'] = 'numericColumn'
        # Sloupce jdou v pořadí, v jakém byly v sešitu → n-tý = písmeno n.
        pismeno = _pismeno(idx + 1)
        if fmt.get(pismeno):
            js = _fmt_js(fmt[pismeno])
            if js:
                c[':valueFormatter'] = js
        if hdr.get(pismeno):
            c['headerStyle'] = hdr[pismeno]
        zaklad = bunky.get(pismeno)
        js = _cf_js(cf[pismeno], zaklad) if cf.get(pismeno) else None
        if js:
            c[':cellStyle'] = js
        elif zaklad:
            c['cellStyle'] = zaklad
        cols.append(c)
    cols += [
        {'headerName': 'Vyjádření', 'field': 'vyjadreni', 'width': 320,
         'editable': editovatelne, 'cellEditor': 'agLargeTextCellEditor',
         'cellEditorPopup': True, 'cellStyle': {'backgroundColor': '#fffbeb'},
         'headerTooltip': 'Vyjádření k řádku monitoru'},
        {'headerName': 'Zapsal', 'field': 'vyjadreni_by', 'width': 150, 'editable': False},
        {'headerName': 'Kdy', 'field': 'vyjadreni_at', 'width': 130, 'editable': False,
         'cellStyle': {'fontSize': '12px', 'color': '#475569'}},
    ]
    return cols


def _export_cols(sloupce: list) -> list:
    cols = [('Období', 'obdobi', 'text', 16)]
    cols += [(s['h'], s['k'], 'num' if s.get('num') else 'text', 16) for s in sloupce]
    cols += [('Vyjádření', 'vyjadreni', 'text', 40),
             ('Zapsal', 'vyjadreni_by', 'text', 18),
             ('Kdy', 'vyjadreni_at', 'text', 16)]
    return cols


# ============================================================
# ==                  VSTUPNÍ OBRAZOVKA                     ==
# ============================================================

@refreshable_na_klienta
async def vykresli_monitor(user_id, user_name: str, vsechna_prava):
    # DB dotazy běží ve vlákně — nedrží event loop celého serveru
    await asyncio.to_thread(inicializace_monitor_db)
    typy = dostupne_typy(vsechna_prava)
    otevreny = app.storage.user.get('monitor_typ')
    if otevreny not in typy:
        otevreny = None

    if otevreny:
        await _vykresli_typ(otevreny, user_id, user_name, vsechna_prava)
        return

    if not typy:
        with ui.column().classes('items-center py-20 gap-3 w-full'):
            ui.icon('lock', size='4rem', color='grey-4')
            ui.label('Nemáte přístup k žádnému monitoru.') \
                .classes('text-lg text-gray-400')
        return

    with ui.row().classes('w-full items-center gap-3 mb-6'):
        ui.icon('monitor_heart', size='2.2rem').classes('text-sky-600')
        ui.label('Monitor').classes('text-3xl font-extrabold text-gray-800')
        ui.label('Sledování cen konkurence – zápis vyjádření a export') \
            .classes('text-sm text-gray-500')

    def _otevri(t):
        async def _fn():
            app.storage.user['monitor_typ'] = t
            # Tamda/makro-ceny mají přes 20 tisíc řádků — než se sestaví grid,
            # ať uživatel vidí, že se něco děje. Překryv musí viset MIMO
            # refreshovaný kontejner, jinak ho překreslení smaže i s dialogem.
            with ui.context.client.layout:
                dlg, _kruh, _popisek = prekryv_kolecko('Načítám přehled…',
                                                       procenta=False)
            dlg.open()
            try:
                await vykresli_monitor.refresh()
            finally:
                dlg.close()
                dlg.delete()
        return _fn

    with ui.row().classes('w-full gap-8 flex-wrap pt-4'):
        for t in typy:
            spec = MONITORY[t]
            intranet_sankce._dlazdice(spec['emoji'], spec['nazev'], spec['border'],
                                      spec['btn'], _otevri(t))


@refreshable_na_klienta
async def _vykresli_typ(typ: str, user_id, user_name: str, vsechna_prava):
    spec = MONITORY[typ]
    tabulka = spec['tabulka']
    _, psat_vse, je_admin = prava_typu(typ, vsechna_prava)

    meta = await asyncio.to_thread(_nacti_meta, tabulka)
    vsechny_styly = await asyncio.to_thread(_nacti_styly, tabulka)
    listy = [l for l in spec['listy'] if l in meta] or list(meta.keys())

    def _zpet():
        app.storage.user['monitor_typ'] = None
        vykresli_monitor.refresh()

    with ui.row().classes('w-full items-center gap-3 mb-6'):
        ui.button(icon='arrow_back', on_click=_zpet).props('flat round color=grey-7') \
            .tooltip('Zpět na přehled monitorů')
        ui.icon('monitor_heart', size='2.2rem').classes('text-sky-600')
        ui.label(f'Monitor {spec["nazev"]}').classes('text-3xl font-extrabold text-gray-800')

    if not listy:
        with ui.column().classes('items-center py-16 gap-3 w-full'):
            ui.icon('inventory_2', size='4rem', color='grey-4')
            ui.label('Zatím nejsou naimportována žádná data.') \
                .classes('text-xl text-gray-400 font-bold')
            if je_admin:
                ui.label('Nahrajte soubor tlačítkem „Nahrát data" vpravo nahoře.') \
                    .classes('text-sm text-gray-400')
                with ui.row().classes('w-full justify-center mt-2'):
                    _import_button(typ, user_name, _vykresli_typ.refresh)
        return

    ulozeny_list = app.storage.user.get(f'monitor_{typ}_list')
    aktualni_list = ulozeny_list if ulozeny_list in listy else listy[0]
    sloupce = meta.get(aktualni_list, [])

    vsechny = await asyncio.to_thread(_nacti, tabulka, aktualni_list)
    # Ořez na sortimenty uživatele. Jediné místo — grid, export i počty
    # čtou přes `_zobrazene()`, takže se nikam nedostane víc, než smí vidět.
    pred_orezem = len(vsechny)
    vsechny = filtruj_dle_sortimentu(vsechny, sloupce, vsechna_prava, je_admin)
    if pred_orezem and not vsechny and not je_admin:
        with ui.column().classes('items-center py-16 gap-3 w-full'):
            ui.icon('lock', size='4rem', color='grey-4')
            ui.label('Nemáte přiřazený žádný nákupní sortiment.') \
                .classes('text-xl text-gray-400 font-bold')
            ui.label('O přiřazení požádejte správce portálu — ve správě '
                     'uživatelů se práva jmenují „Sortiment: …".') \
                .classes('text-sm text-gray-400')
        return
    obdobi_list = _seznam_obdobi(tabulka, aktualni_list)
    ulozene_obd = app.storage.user.get(f'monitor_{typ}_obdobi')
    stav = {'obdobi': ulozene_obd if ulozene_obd in obdobi_list
                      else (obdobi_list[0] if obdobi_list else None),
            'jen_bez': False}

    def _zobrazene():
        data = vsechny
        if stav['obdobi']:
            data = [r for r in data if r.get('obdobi') == stav['obdobi']]
        if stav['jen_bez']:
            data = [r for r in data if not (r.get('vyjadreni') or '').strip()]
        return data

    def _info_text(data=None):
        data = _zobrazene() if data is None else data
        bez = sum(1 for r in data if not (r.get('vyjadreni') or '').strip())
        return f'Zobrazeno řádků: {len(data)} · bez vyjádření: {bez}'

    async def _export():
        data = _zobrazene()
        ids = await intranet_sankce._viditelne_ids(grid)
        if ids:
            data, _ = intranet_sankce._serad_dle_ids(data, ids)
        if not data:
            ui.notify('Aktuální filtr nevrací žádné řádky — není co exportovat.',
                      type='warning', position='top', timeout=6000)
            return
        zaklad = f'monitor_{typ}_{intranet_sankce._safe_filename(aktualni_list)}'
        sablona, pripona = await asyncio.to_thread(_nacti_sablonu, tabulka)
        if sablona:
            try:
                bajty = await asyncio.to_thread(
                    _do_sablony, sablona, pripona, aktualni_list, sloupce, data,
                    [('Vyjádření', 'vyjadreni'), ('Zapsal', 'vyjadreni_by'),
                     ('Kdy', 'vyjadreni_at')])
                jmeno = f'{zaklad}_{datetime.datetime.now():%Y-%m-%d_%H%M}.{pripona}'
                intranet_sankce._stahni_pres_http(bajty, jmeno)
                return
            except Exception as e:
                print(f'[monitory] Export ze šablony selhal ({e}) — jedu bez formátů.')
        await _export_xlsx(_export_cols(sloupce), data, None, zaklad, spec['nazev'])

    def _aplikuj():
        data = _zobrazene()
        grid.options['rowData'] = data
        grid.update()
        info.set_text(_info_text(data))

    # -- Ovládací lišta --
    with ui.row().classes('w-full items-center gap-3 mb-2 flex-wrap'):
        if len(listy) > 1:
            def _on_list(e):
                app.storage.user[f'monitor_{typ}_list'] = e.value
                _vykresli_typ.refresh()
            ui.select(listy, value=aktualni_list, label='List') \
                .props('outlined dense options-dense').classes('w-56') \
                .on_value_change(_on_list) \
                .tooltip('Sešit monitoru obsahuje víc sestav — vyberte, kterou chcete vidět.')

        def _on_obd(e):
            stav['obdobi'] = None if e.value == _VSE_OBD else e.value
            app.storage.user[f'monitor_{typ}_obdobi'] = stav['obdobi']
            _aplikuj()
        ui.select([_VSE_OBD] + obdobi_list, value=stav['obdobi'] or _VSE_OBD,
                  label='Období').props('outlined dense options-dense').classes('w-60') \
            .on_value_change(_on_obd)

        def _on_bez(e):
            stav['jen_bez'] = bool(e.value)
            _aplikuj()
        ui.switch('Jen bez vyjádření', value=False).on_value_change(_on_bez) \
            .tooltip('Zobrazí pouze řádky, ke kterým zatím nikdo nic nenapsal.')

        ui.space()
        info = ui.label(_info_text()).classes('text-sm text-gray-500')
        ui.button(icon='download', text='Export', on_click=_export) \
            .props('color=secondary outline dense no-caps') \
            .tooltip('Stáhne .xlsx podle aktuálního filtru (včetně filtrů v hlavičkách).')
        if je_admin:
            _import_button(typ, user_name, _vykresli_typ.refresh)
            _smazat_button(typ, user_name, _vykresli_typ.refresh)

    if psat_vse:
        ui.label('✍️ Vyjádření píšete přímo do žlutého sloupce — zápis se propíše '
                 'do všech označených řádků.').classes('text-xs text-gray-500 mb-1')
    else:
        ui.label('👁️ Sestavu vidíte jen pro čtení – filtrovat, řadit a exportovat '
                 'ale můžete.').classes('text-xs text-gray-500 mb-1')

    grid = ui.aggrid({
        'columnDefs': _col_defs(sloupce, psat_vse, vsechny_styly.get(aktualni_list)),
        'rowData': _zobrazene(),
        'defaultColDef': {'resizable': True, 'sortable': False, 'filter': True},
        'rowHeight': 32,
        'singleClickEdit': True,
        'stopEditingWhenCellsLoseFocus': True,
        'suppressMovableColumns': True,
        'rowSelection': 'multiple',
        'suppressRowClickSelection': True,
        ':getRowId': intranet_sankce._GET_ROW_ID,
        ':onFirstDataRendered': intranet_sankce._AUTOSIZE_FIT,
    }).classes('w-full').style(intranet_sankce._GRID_STYLE)

    # -- Zápis vyjádření --
    async def _on_change(e):
        a = e.args or {}
        if a.get('colId') != 'vyjadreni':
            return
        d = a.get('data') or {}
        rid = d.get('id')
        nova = a.get('newValue')
        radek = next((r for r in vsechny if r.get('id') == rid), None)
        if radek is None:
            return
        # Zápis se propíše do všech právě označených řádků (jinak jen do toho jednoho).
        cile = [radek]
        try:
            sel_ids = await ui.run_javascript(
                f"const c=getElement({grid.id});"
                "return (c&&c.api&&c.api.getSelectedRows)?"
                "c.api.getSelectedRows().map(r=>r.id):[];", timeout=5)
        except Exception:
            sel_ids = []
        if sel_ids and rid in sel_ids:
            cile = [r for r in vsechny if r.get('id') in set(sel_ids)]
        pocet = _uloz_vyjadreni(tabulka, cile, _s(nova or ''), user_id, user_name)
        if not pocet:
            return
        intranet_logger.log_activity(
            user_name, 'Monitor',
            f'Vyjádření {spec["nazev"]} – {pocet} řádků')
        if pocet > 1:
            ui.notify(f'Vyjádření zapsáno u {pocet} řádků.', type='positive',
                      position='top', timeout=4000)
        grid.run_grid_method('applyTransaction', {'update': cile})
        info.set_text(_info_text())
    grid.on('cellValueChanged', _on_change)

    # -- Klik na očičko: historie řádku --
    def _on_click(e):
        a = e.args or {}
        if a.get('colId') != '_eye':
            return
        d = a.get('data') or {}
        popis = ' – '.join(x for x in ((d.get(s['k']) and _s(d.get(s['k'])))
                                       for s in sloupce[:2]) if x)
        _zobraz_historii(tabulka, d.get('row_hash'), popis)
    grid.on('cellClicked', _on_click)


# ============================================================
# ==                  IMPORT / MAZÁNÍ (správce)             ==
# ============================================================

def _otevri_import_dialog(typ: str, user_name: str, refresh_fn):
    spec = MONITORY[typ]
    drzeny = {'raw': None, 'name': ''}
    with ui.dialog() as dlg, ui.card().classes('p-6 rounded-2xl gap-3') \
            .style('min-width: 480px; max-width: 560px'):
        with ui.row().classes('items-center gap-2'):
            ui.icon('upload_file', color='primary').classes('text-2xl')
            ui.label(f'Nahrát data – Monitor {spec["nazev"]}') \
                .classes('text-xl font-bold text-gray-800')
        ui.label('Vyberte období, ke kterému data patří. Dávka stejného období se '
                 'nahradí, dříve zapsaná vyjádření zůstanou zachována.') \
            .classes('text-xs text-gray-500 -mt-1')
        ui.label('Listy v sešitu: ' + ', '.join(f'„{l}"' for l in spec['listy'])) \
            .classes('text-xs text-gray-400')

        with ui.row().classes('items-end gap-3 w-full') as row_datum:
            inp_od = ui.input('Období OD').props('type=date outlined dense').classes('flex-1')
            inp_do = ui.input('Období DO').props('type=date outlined dense').classes('flex-1')
        row_datum.set_visibility(True)
        hint = ui.label('Necháte-li prázdné, zkusím období načíst z názvu souboru.') \
            .classes('text-xs text-gray-400 -mt-1')

        stav_lbl = ui.label('').classes('text-sm text-gray-600')
        stav_lbl.set_visibility(False)

        async def _zprac(raw: bytes, name: str):
            od_iso = (inp_od.value or '').strip()
            do_iso = (inp_do.value or '').strip()
            if not (od_iso and do_iso):
                p_od, p_do = _obdobi_z_nazvu(name)
                od_iso = od_iso or p_od
                do_iso = do_iso or p_do
            drzeny['raw'] = raw
            drzeny['name'] = name
            if not (od_iso and do_iso):
                stav_lbl.set_text(f'Soubor „{name}" je načten. Doplňte období OD i DO '
                                  'a klikněte na „Importovat".')
                stav_lbl.set_visibility(True)
                btn_import.set_visibility(True)
                return
            await _spust(raw, name, od_iso, do_iso)

        async def _spust(raw: bytes, name: str, od_iso: str, do_iso: str):
            obdobi = _obdobi_label(od_iso, do_iso)
            stav_pokrok = {'f': 0.0}
            dlg.close()
            kol_dlg, kruh, kol_popisek = prekryv_kolecko(f'Importuji {name}…')
            kol_dlg.open()

            def _tik():
                kruh.set_value(stav_pokrok['f'] * 100)
                kol_popisek.set_text(f'{stav_pokrok["f"] * 100:.0f} %')
            casovac = ui.timer(0.2, _tik)
            try:
                count, listy, err = await asyncio.to_thread(
                    _importuj_sync, raw, typ, obdobi, od_iso, do_iso, user_name,
                    lambda f: stav_pokrok.update(f=f), name)
            finally:
                casovac.cancel()
                kol_dlg.close()
            if err:
                ui.notify(f'Import se nezdařil: {err}', type='negative', timeout=10000)
                dlg.open()
                return
            intranet_logger.log_activity(
                user_name, 'Monitor',
                f'Import {spec["nazev"]} ({", ".join(listy)}) – období {obdobi}: '
                f'{count} řádků')
            ui.notify(f'Import dokončen — období {obdobi}, {count} řádků '
                      f'({", ".join(listy)}).', type='positive', position='top-right',
                      timeout=6000)
            drzeny['raw'] = None
            refresh_fn()

        async def _on_upload(e):
            # Napříč verzemi NiceGUI se obsah nahraného souboru jmenuje různě.
            zdroj = None
            for attr in ('content', 'file', 'stream', 'data', 'file_obj'):
                val = getattr(e, attr, None)
                if val is not None and hasattr(val, 'read'):
                    zdroj = val
                    break
            if zdroj is None:
                ui.notify('Nepodařilo se načíst obsah souboru.', type='negative')
                up.reset()
                return
            try:
                raw = zdroj.read()
                if inspect.isawaitable(raw):
                    raw = await raw
            except Exception as exc:
                ui.notify(f'Chyba čtení souboru: {exc}', type='negative')
                up.reset()
                return
            # Bytes si držíme sami; widget hned uvolníme, ať jde příště nahrát znovu.
            up.reset()
            # NiceGUI 2.x mělo jméno na události, 3.x na FileUpload (e.file.name).
            # Bez jména se neurčí přípona → nevytáhne se šablona a padá i období.
            jmeno = getattr(e, 'name', '') or getattr(zdroj, 'name', '') or drzeny['name']
            await _zprac(raw, jmeno)

        up = ui.upload(on_upload=_on_upload, auto_upload=True, max_file_size=50_000_000,
                       label='Vybrat soubor .xlsx / .xlsm / .xlsb') \
            .props('accept=.xlsx,.xlsm,.xlsb').classes('w-full')

        btn_import = ui.button('Importovat', icon='play_arrow', on_click=lambda: _spust(
            drzeny['raw'], drzeny['name'],
            (inp_od.value or '').strip(), (inp_do.value or '').strip())) \
            .props('unelevated no-caps').classes('bg-blue-600 text-white')
        btn_import.set_visibility(False)

        with ui.row().classes('w-full justify-end'):
            ui.button('Zavřít', on_click=dlg.close).props('flat no-caps color=grey-7')
    dlg.open()


def _import_button(typ: str, user_name: str, refresh_fn):
    ui.button('Nahrát data', icon='upload_file',
              on_click=lambda: _otevri_import_dialog(typ, user_name, refresh_fn)) \
        .props('unelevated no-caps') \
        .classes('bg-blue-600 hover:bg-blue-700 text-white font-semibold rounded-lg shadow-md px-5')


def _smazat_button(typ: str, user_name: str, refresh_fn):
    spec = MONITORY[typ]

    def _otevri():
        obdobi_list = _seznam_obdobi(spec['tabulka'])
        VSE = '(všechna období)'
        with ui.dialog() as dlg, ui.card().classes('p-6 rounded-2xl gap-3'):
            with ui.row().classes('items-center gap-2'):
                ui.icon('warning', color='red').classes('text-2xl')
                ui.label(f'Smazat data – Monitor {spec["nazev"]}') \
                    .classes('text-xl font-bold text-gray-800')
            ui.label('Mazání je nevratné a smaže i historii změn (včetně zapsaných '
                     'vyjádření). Vyberte rozsah:').classes('text-sm text-gray-600')
            sel = ui.select([VSE] + obdobi_list, value=obdobi_list[0] if obdobi_list else VSE,
                            label='Období').props('outlined dense options-dense').classes('w-72')

            def _smaz():
                obd = None if sel.value == VSE else sel.value
                pocet, err = _smaz_data(spec['tabulka'], obd)
                dlg.close()
                if err:
                    ui.notify(err, type='negative')
                    return
                intranet_logger.log_activity(
                    user_name, 'Monitor',
                    f'Smazání {spec["nazev"]} – {obd or "vše"}: {pocet} řádků')
                ui.notify(f'Smazáno {pocet} řádků.', type='positive', position='top')
                refresh_fn()

            with ui.row().classes('w-full justify-end gap-2'):
                ui.button('Zrušit', on_click=dlg.close).props('flat no-caps color=grey-7')
                ui.button('Smazat', icon='delete_forever', on_click=_smaz) \
                    .props('unelevated no-caps color=red')
        dlg.open()

    ui.button('Smazat data', icon='delete_forever', on_click=_otevri) \
        .props('outline no-caps color=red dense')
